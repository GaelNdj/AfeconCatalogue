#!/usr/bin/env python3
"""
Extraction catalogue CEDEO — images pièces, netteté native (sans zoom).

- Pixels natifs du PDF (pas de rendu ×3 qui agrandissait ~44 %)
- PNG sans recompression JPEG
- Ignore les photos de personnes / lifestyle
- Marge blanche pour éviter un recadrage trop serré

Usage:
  python3 scripts/extract_catalog_hq.py
"""

import hashlib
import math
import os
import re
import sys
from collections import Counter

import fitz
import openpyxl
import pypdfium2 as pdfium
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

PDF_PATH = "/Users/gael/Downloads/cata_pro_2026_idf_368_enrich.pdf"
CATALOGUE_DIR = "/Users/gael/Desktop/catalogue"
OUTPUT_PATH = os.path.join(CATALOGUE_DIR, "catalogue_pro_2026_plomberie.xlsx")
if os.environ.get("EXTRACT_OUTPUT"):
    OUTPUT_PATH = os.environ["EXTRACT_OUTPUT"]
IMAGE_DIR = os.path.join(CATALOGUE_DIR, "images_hq")
IMAGE_REL = "images_hq"
if os.environ.get("EXTRACT_IMAGE_DIR"):
    IMAGE_DIR = os.environ["EXTRACT_IMAGE_DIR"]
    IMAGE_REL = os.path.basename(IMAGE_DIR.rstrip("/"))
FILTER_FAMILLE = "Plomberie"
MAX_NAME_LEN = 90
# Native pixels only — no page render upsample (that zoomed photos ~1.44x).
PAD_FRAC = 0.08
MAX_IMAGE_PX = 280
IMAGE_EXT = "png"
LETTER_PREFIX_RE = re.compile(r"^([A-Da-d])\s+")
SECTION_PHOTO_RE = re.compile(r"^●\s*(DROITE|COUDÉE|COUDEE)", re.I)

COLUMNS = [
    "Réf.Pro",
    "Réf.Four",
    "Diamètre",
    "Vendu par",
    "Marque",
    "Désignation",
    "Variante",
    "Ordre",
    "Code",
    "Code interne",
    "Famille",
    "Catégorie",
    "Prix HT",
    "Note",
    "Description",
    "Image",
    "Image produit",
]

SKIP_LINE = (
    "COMMANDEZ EN LIGNE",
    "DISPONIBLE",
    "SOUS 24H",
    "SOUS 3 JOURS",
    "SUR COMMANDE",
    "ECO-CONTRIBUTION",
    "PRIX INDICATIF",
    "RÉVISABLE",
    "ARTICLE TOP PRIX",
    "BAISSE",
    "NOSPRIX",
    "CEDEO.FR",
    "RETOUVEZ TOUTE",
    "UNE QUESTION",
    "VOS OFFRES",
    "PROMOS CEDEO",
    "- GARANTIE",
)

BRANDS = [
    "JACOB DELAFON", "IDEAL STANDARD", "HANSGROHE", "SAUNIER DUVAL",
    "DE DIETRICH", "ELM LEBLANC", "SYNTHÈSE MINÉRAL", "CONEX",
    "ALTECH", "WIELAND", "TALOS", "STARFIX", "SANCO", "CUPROLIFE",
    "WICU", "REHAU", "MICROFLEX", "GROHE", "PORCHER", "GEBERIT",
    "DURAVIT", "VILLEROY", "ROCA", "ACQUABELLA", "KINEDO", "SANITRIT",
    "SFA", "ATLANTIC", "THERMOR", "CHAPPEE", "VIESSMANN", "BOSCH",
    "DAIKIN", "MITSUBISHI", "HITACHI", "TOSHIBA", "PANASONIC",
    "AIRWELL", "FRISQUET", "NOVELLINI", "WIRQUIN", "VALENTIN",
    "PRESTO", "DELABIE", "SCHELL", "WATTS", "HONEYWELL", "DANFOSS",
    "CALEFFI", "GIACOMINI", "COMAP", "SOMATHERM",     "NICOLL", "GEBO",
    "NOVIPRO", "KALDEWEI", "SIAMP", "SYNTHÈSE", "TECHNO",
    "FINIMETAL", "ACOVA", "ZEHNDER", "AYOR", "ALTERNA",
    "VIRAX", "ROTHENBERGER", "KNIPEX", "GRUNDFOS", "BANIDES",
    "WILO", "REMS", "SUPER-EGO", "RIDGID",
]

SERIES_BRAND = {
    "BANGA": "Finimetal",
    "BANGA RACCORDEMENT CENTRAL": "Finimetal",
    "BANGA CINTRE ÉLECTRIQUE": "Finimetal",
    "BANGA ÉLECTRIQUE": "Finimetal",
    "KIT FIXOPLAC": "Ayor",
}

TABLE_SECTION_NAMES = {
    "UNI", "BICOLORE", "VERSION DROITE", "VERSION GAUCHE",
    "PLAN VASQUE", "MEUBLE", "CONSOLE", "ACCESSOIRES",
}
TECH_VARIANT_RE = re.compile(
    r"^A\s+(GLISSEMENT|SERTIR|SERTISSAGE|COMPRESSION|COLLER)\b",
    re.I,
)

SLOGAN_FRAGMENTS = (
    "15 MIN DE VOS CHANTIERS",
    "3000 EXPERTS",
    "À VOS CLIENTS",
    "À VOS CÔTÉS",
    "UNE AGENCE À MOINS",
    "UNE AVANCE DE TRÉSORERIE",
    "TOUJOURS UNE AGENCE",
    "PROCHE DE MON CHANTIER",
    "NORMES ÉLECTRIQUES",
    "ÉLECTRIQUES APPLIQUÉES",
    "À LA SALLE DE BAIN",
)

TITLE_STOPWORDS = {
    "ANS", "ON", "HORS", "LE", "DE", "PRO", "DES", "LES", "UNE", "AUX",
    "PAR", "SUR", "THE", "VOLUME", "CONSEIL", "CENTRAL", "BLANC", "NOIR",
    "UNI", "NOUVEAU",
}

TITLE_CONTINUE_WORDS = {
    "CENTRAL", "DROITE", "GAUCHE", "ÉLECTRIQUE", "ELECTRIQUE",
    "SOUFFLANT", "MIXTE", "ASYMÉTRIQUE", "SYMETRIQUE", "SYMÉTRIQUE",
}

FAMILLE_RULES = [
    (["OUTILLAGE", "EPI", "NOVIPRO", "FRIGORISTE"], "Outillage"),
    (["CARRELAGE", "REVÊTEMENT", "REVETEMENT"], "Revêtements"),
    (
        [
            "CHAUDI", "P.A.C", "PAC ", "AIR/EAU", "AIR/AIR", "RADIATEUR",
            "SÈCHE-SERVIETTE", "SECHE-SERVIETTE",
            "CLIM", "VENTIL", "CHAUFFE", "FUMIST", "CHAUFF", "GÉNIE CLIM",
            "GENIE CLIM", "PLANCHE", "THERMOSTAT", "POMPE À CHALEUR",
            "POMPES À CHALEUR",
        ],
        # Dans le site, le génie climatique est une sous-famille de Plomberie,
        # pas une famille à part : sinon MAGNA / À BRIDES sortent de l'Excel.
        "Plomberie",
    ),
    (
        [
            "RACCORD", "TUBE", "CUIVRE", "LAITON", "PER", "MULTICOUCHE",
            "PVC", "FONTE", "ACIER", "INOX", "PLOMBERIE", "SERTIR",
            "ELECTROZINGU",
        ],
        "Plomberie",
    ),
    (
        [
            "RECEVEUR", "PAROI", "DOUCHE", "BAIN", "WC", "LAVABO",
            "SANITAIRE", "ROBINET", "MITIGEUR", "ABATTANT", "MEUBLE",
            "VIDAGE", "SIPHON", "BONDE",
        ],
        "Sanitaire",
    ),
]

CODE_RE = re.compile(r"\b(\d{7})\b")
PRICE_RE = re.compile(r"(\d{1,5}[.,]\d{2})\s*$")
PRICE_COURS_RE = re.compile(r"prix\s*[àa]?\s*cours", re.I)
DIAM_RE = re.compile(
    r"^(\d{1,2}(?:[-/]\d{1,2})?(?:[xX×]\d{1,2}){0,2}(?:-\d{1,2})?)$"
)
SIZE_RE = re.compile(
    r"(\d{2,4}\s*[xX×]\s*\d{2,4}(?:\s*[xX×]\s*\d{2,4})?\s*(?:cm|mm)?|\d{2,4}\s*cm)\b",
    re.I,
)
BULLET_NOTE_RE = re.compile(r"^[●•▪]\s*(.+)$")
# Triangle catalogue (●/◗) ; fitz le transcrit souvent en « Q ».
TABLE_BULLET_RE = re.compile(r"^(?:[●•▪◗▸►]\s*|Q\s+)")
# Les puces du catalogue sont des glyphes ZapfDingbats que fitz rend comme des
# lettres (« G », « Q »…) : on les repère par la police, pas par le caractère.
SYMBOL_FONTS = ("ZapfDingbats", "Dingbat", "Wingdings", "Symbol")
HEADER_ROW_RE = re.compile(
    r"(RéfPro|Réf\.?\s*Four|Prix HT|Vendu par|Désignation|Poids)",
    re.I,
)
ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
WRAP_END_RE = re.compile(r"\d+\s*[xX×]\s*\d*\s*$")


def clean_spaces(s):
    return re.sub(r"\s+", " ", (s or "").replace("\xa0", " ")).strip()


def sanitize_cell(val):
    if val is None:
        return ""
    s = ILLEGAL_XLSX.sub(" ", str(val))
    return clean_spaces(s)


def is_category_header(line, current_category=""):
    """Rubriques catalogue (PLOMBERIE, RACCORDS CUIVRE…) ≠ fiches produit."""
    s = clean_spaces(line)
    if not s or s[:1].islower():
        return False
    if is_slogan_line(s):
        return False
    if " - " in s:
        return False
    up = s.upper().rstrip(".")
    if current_category and up == clean_spaces(current_category).upper():
        return True
    headers = {
        "PLOMBERIE", "SANITAIRE", "OUTILLAGE", "GÉNIE CLIMATIQUE", "GENIE CLIMATIQUE",
        "RACCORDS CUIVRE", "RACCORDS LAITON", "TUBES CUIVRE", "TUBES ET RACCORDS CUIVRE ET LAITON",
        "TUBES ET RACCORDS FONTE", "GAINÉS ISOLÉS", "SÈCHE-SERVIETTES", "COLLECTEURS",
        "TUBES ET RACCORDS", "RACCORDS", "TUBES", "PIÈCES", "PIECES", "ACCESSOIRES",
    }
    if up in headers:
        return True
    if up.startswith("TUBES ET RACCORDS") and "Ø" not in up and "°" not in up:
        return True
    if re.fullmatch(r"(RACCORDS|TUBES|ROBINETS|COLLECTEURS|VANNES)\s+[A-ZÀ-Ÿ]+$", up):
        return True
    return False


def is_warranty_or_meta_line(line):
    """GARANTIE / promos = méta, jamais titre de fiche produit."""
    s = clean_spaces(line)
    up = s.upper()
    if re.fullmatch(r"GARANTIE\.?", up):
        return True
    if up.startswith("GARANTIE :") or up.startswith("- GARANTIE"):
        return True
    if "VOS OFFRES" in up or "PROMOS CEDEO" in up:
        return True
    return False


def line_column(tp, line, page_w, page_h):
    """Colonne L/R d'une ligne (titre ou ligne code) via position dans le PDF."""
    if CODE_RE.search(line):
        code = CODE_RE.findall(line)[-1]
        y, x = find_title_y(tp, code, page_h)
        if x is not None:
            return "L" if x < page_w / 2 else "R"
    for needle in (line[:32], line[:20], line[:12]):
        if len(needle) < 4:
            continue
        y, x = find_title_y(tp, needle, page_h)
        if x is not None:
            return "L" if x < page_w / 2 else "R"
    return "L"


def line_text_with_bullets(spans):
    """Texte d'une ligne, puces dingbats rétablies en « ● »."""
    parts = []
    for span in spans:
        txt = span.get("text", "")
        font = span.get("font") or ""
        if len(txt.strip()) == 1 and any(f in font for f in SYMBOL_FONTS):
            parts.append("● ")
        else:
            parts.append(txt)
    return clean_spaces("".join(parts))


# Le catalogue suit une hiérarchie typographique stricte : le titre d'article est
# le seul texte en Montserrat-SemiBold 8.11 et la description courte le seul en
# Montserrat-Light 6.50. Les cellules de tableau sont en 5.50. C'est bien plus
# fiable que de deviner d'après le contenu : sans ça, la fin d'une description ou
# d'une cellule qui passe à la ligne finit prise pour un nom d'article.
TITLE_FONT, TITLE_SIZE = "Montserrat-SemiBold", 8.11
DESC_FONT, DESC_SIZE = "Montserrat-Light", 6.50


def line_role(spans):
    """« title », « desc », ou "" — d'après la police du premier span visible."""
    for span in spans:
        if not (span.get("text") or "").strip():
            continue
        font = span.get("font") or ""
        size = float(span.get("size") or 0)
        if font == TITLE_FONT and abs(size - TITLE_SIZE) < 0.15:
            return "title"
        if font == DESC_FONT and abs(size - DESC_SIZE) < 0.15:
            return "desc"
        return ""
    return ""


def extract_column_items(fitz_page, page_w):
    """Lignes (y, x0, texte, rôle) par colonne, haut → bas. Ignore l'onglet latéral."""
    by_col = {"L": [], "R": []}
    for block in fitz_page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            text = line_text_with_bullets(line.get("spans", []))
            if not text:
                continue
            bbox = line.get("bbox", (0, 0, 0, 0))
            x0, y0, x1, _y1 = bbox
            if float(x0) >= page_w * 0.90:
                continue
            cx = (float(x0) + float(x1)) / 2
            col = "L" if cx < page_w / 2 else "R"
            by_col[col].append(
                (float(y0), float(x0), text, line_role(line.get("spans", [])))
            )
    for col in by_col:
        by_col[col].sort(key=lambda x: (x[0], x[1]))
    return by_col


def merge_wrapped_items(items):
    out = []
    i = 0
    n = len(items)
    while i < n:
        y, line = items[i]
        if not CODE_RE.search(line) and i + 1 < n:
            ny, nxt = items[i + 1]
            join = False
            if ny - y < 16 and WRAP_END_RE.search(line):
                join = True
            elif ny - y < 14 and re.match(r"^(cm|mm)\b", nxt, re.I) and re.search(r"\d", line):
                join = True
            elif (
                ny - y < 14
                and looks_like_description_line(line)
                and looks_like_description_line(nxt)
            ):
                join = True
            if join:
                combined = clean_spaces(line + " " + nxt)
                i += 1
                if not CODE_RE.search(combined) and i + 1 < n and CODE_RE.search(items[i + 1][1]):
                    combined = clean_spaces(combined + " " + items[i + 1][1])
                    i += 1
                out.append((y, combined))
                i += 1
                continue
        out.append((y, line))
        i += 1
    return out


def strip_table_bullet(line):
    s = TABLE_BULLET_RE.sub("", clean_spaces(line)).strip()
    m = re.match(r"^([A-Z])\s+(.+)$", s)
    if m and (
        m.group(2).upper() in TABLE_SECTION_NAMES or TECH_VARIANT_RE.match(m.group(2))
    ):
        return m.group(2).strip()
    return s


def is_table_section_header(line):
    """Sous-rubrique de tableau (A GLISSEMENT, BAIN/DOUCHE…) ≠ nom d'article."""
    s = clean_spaces(line)
    if not s or CODE_RE.search(s) or PRICE_RE.search(s) or HEADER_ROW_RE.search(s):
        return False
    rest = strip_table_bullet(s)
    reason = None
    if looks_like_row_prefix(s) or looks_like_row_prefix(rest):
        return False
    if DIAM_RE.match((rest or "").replace(" ", "")) or DIAM_RE.match(s.replace(" ", "")):
        return False
    if TABLE_BULLET_RE.match(s) or rest != s:
        if re.fullmatch(r"[A-Da-d]", rest or ""):
            return False
        if rest.upper() in TABLE_SECTION_NAMES or TECH_VARIANT_RE.match(rest):
            reason = "bullet_named_or_tech"
        elif TABLE_BULLET_RE.match(s) and 3 <= len(rest) <= 90:
            reason = "bullet_any_rest"
        if reason:
            return True
    if TECH_VARIANT_RE.match(s) or TECH_VARIANT_RE.match(rest):
        return True
    if rest.upper() in TABLE_SECTION_NAMES or s.upper() in TABLE_SECTION_NAMES:
        return True
    if re.search(r"\s/\s", s) and re.search(
        r"MODULE|BAIN|DOUCHE|SOUS-ÉVIER|SOUS-EVIER|LAVABO|WC\b", s, re.I
    ):
        # Vrais en-têtes de tableau (« BAIN / DOUCHE »), pas un nom de 80 caractères
        # qui contient un slash et le mot douche (ex. CROMA SELECT / ECOSTAT).
        if len(s) <= 40:
            return True
    return False


def apply_table_section(pend, line):
    label = strip_table_bullet(line) or clean_spaces(line)
    up = label.upper()
    if TECH_VARIANT_RE.match(label):
        pend["variant_sub"] = label
    elif up in {"UNI", "BICOLORE"} and pend.get("variant_group"):
        pend["variant_sub"] = label
    else:
        pend["variant_group"] = label
        pend["variant_sub"] = ""


def band_end_for(y, seps):
    """Y du trait vert/rouge qui ferme l'article contenant y."""
    if y is None or not seps or len(seps) < 2:
        return None
    for i in range(len(seps) - 1):
        if seps[i] - 10 <= y < seps[i + 1]:
            return seps[i + 1]
    return seps[-1]


def is_layout_chrome(line):
    s = clean_spaces(line)
    if not s:
        return True
    if re.fullmatch(r"[A-Da-d]", s):
        return True
    if s in {"Ø", "(mm)", "mm", "RéfPro", "Code", "Prix HT"}:
        return True
    if HEADER_ROW_RE.search(s) and not CODE_RE.search(s):
        return True
    return False


ROW_PREFIX_RE = re.compile(
    r"^(?:[a-dA-D●•]\s+)?Ø?\s*\d{1,2}(?:[.,]\d+)?"
    r"(?:[-/]\d{1,2})?(?:[xX×]\d{1,2}){0,2}(?:-\d{1,2})?"
    r"(?:\s+\d{3,5})?$"
)


def looks_like_row_prefix(line):
    """Ligne Ø / réf. pro, avant le code 7 chiffres de la ligne suivante."""
    s = clean_spaces(line)
    if not s or CODE_RE.search(s) or PRICE_RE.search(s):
        return False
    return bool(ROW_PREFIX_RE.match(s))


def new_pending(category="", brand=""):
    return {
        "category": category,
        "title": "",
        "brand": brand,
        "note": "",
        "description": "",
        "vendu": "",
        "image": "",
        "variant_group": "",
        "variant_sub": "",
        "_section_letter": "",
        "_pending_letter": "",
        "_band_end": None,
    }


def is_skip_line(line):
    up = line.upper()
    return any(k in up for k in SKIP_LINE)


def normalize_price(raw):
    if not raw:
        return ""
    if PRICE_COURS_RE.search(raw):
        return "Prix à cours"
    raw = raw.replace(" ", "").replace(".", ",")
    m = re.match(r"^(\d+),(\d{2})$", raw)
    if not m:
        return ""
    val = float(f"{m.group(1)}.{m.group(2)}")
    if val == 0.01 or val <= 0 or val > 20000:
        return ""
    return f"{m.group(1)},{m.group(2)}"


def detect_brand(text):
    up = text.upper()
    for b in BRANDS:
        if b in up:
            if b in ("SYNTHÈSE", "SYNTHÈSE MINÉRAL"):
                return "SYNTHÈSE MINÉRAL"
            if b == "ALTECH":
                return "Altech"
            return b.title() if b != b.upper() else b
    return ""


def infer_famille(category, title):
    blob = f"{category} {title}".upper()
    # Outillage du plombier : reste dans le catalogue Plomberie, pas une famille à part.
    if "PLOMBIER" in blob:
        return "Plomberie"
    for keys, famille in FAMILLE_RULES:
        if any(k in blob for k in keys):
            return famille
    return "Autres"


def is_slogan_line(line):
    up = clean_spaces(line).upper()
    if not up:
        return False
    return any(s in up for s in SLOGAN_FRAGMENTS)


def looks_like_description_line(line):
    s = clean_spaces(line)
    if not s or len(s) < 12:
        return False
    if CODE_RE.search(s) or PRICE_RE.search(s) or HEADER_ROW_RE.search(s):
        return False
    if is_skip_line(s) or is_warranty_or_meta_line(s) or is_slogan_line(s):
        return False
    if is_category_header(s):
        return False
    letters = [c for c in s if c.isalpha()]
    if not letters:
        return False
    lower_ratio = sum(c.islower() for c in letters) / len(letters)
    return lower_ratio > 0.25 or ":" in s


def _soften_desc_clause(clause):
    """Légère réécriture d'une phrase du catalogue, sans changer le sens."""
    c = clean_spaces(clause or "").strip(" -–—.;")
    if not c:
        return ""
    # « Pour la circulation de liquide… » → « Circulation de liquide… »
    c = re.sub(
        r"^(?:pour|permet(?:tant)?(?:\s+de)?|destiné[e]?\s+à|sert\s+à)\s+(?:la |le |les |l['’])?",
        "",
        c,
        flags=re.I,
    )
    c = re.sub(r"\bafin de\b", "pour", c, flags=re.I)
    c = re.sub(r"\bselon les besoins de l['’]installation\b", "", c, flags=re.I)
    c = re.sub(r"\s{2,}", " ", c).strip(" ,;.")
    inf = {
        "optimiser": "Optimise",
        "raccorder": "Raccorde",
        "assurer": "Assure",
        "permettre": "Permet",
        "faciliter": "Facilite",
        "réduire": "Réduit",
        "reduire": "Réduit",
        "adapter": "Adapte",
    }
    first = c.split(" ", 1)
    mapped = inf.get(first[0].casefold())
    if mapped:
        c = mapped if len(first) == 1 else mapped + " " + first[1]
    if c:
        c = c[0].upper() + c[1:]
    return c


def rewrite_short_description(raw, title=""):
    """Paragraphe Light du catalogue → description courte, un peu plus concise."""
    t = clean_spaces(raw or "")
    t = re.sub(r"[\x00-\x1f]", "", t)
    # Ligatures PDF coupées par une espace : « aﬁ n », « modiﬁ er ».
    t = re.sub(r"ﬁ (?=[a-zà-ÿ])", "fi", t)
    t = re.sub(r"ﬂ (?=[a-zà-ÿ])", "fl", t)
    t = t.replace("ﬁ", "fi").replace("ﬂ", "fl").replace("œ", "oe")
    t = re.sub(r"\bfi letage\b", "filetage", t, flags=re.I)
    t = re.sub(r"\bfl uide\b", "fluide", t, flags=re.I)
    t = re.sub(r"\s*J\s*[\d,.]+\s*", " ", t)
    t = re.sub(r"\s+", " ", t).strip(" -–—")
    if not t:
        return None
    t = re.sub(r"^Raccordements\s*:\s*", "Raccordement : ", t, flags=re.I)
    t = re.sub(
        r"\s*[-–—]\s*Pression de service\s*:\s*",
        ". Pression de service : ",
        t,
        flags=re.I,
    )
    # Ne pas recopier le titre : ça remplaçait la vraie description.
    title_cmp = clean_spaces(title or "").casefold()
    if title_cmp and t.casefold() == title_cmp:
        return None
    parts = [
        p.strip(" -–—.;:")
        for p in re.split(r"\s+[-–—]\s+|(?<=[.;])\s+|:\s+", t)
        if p.strip(" -–—.;:")
    ]
    clauses = []
    for p in parts:
        softened = _soften_desc_clause(p)
        if not softened:
            continue
        if title_cmp and softened.casefold() == title_cmp:
            continue
        if any(softened.casefold() == c.casefold() for c in clauses):
            continue
        clauses.append(softened)
    if not clauses:
        return None
    # Deux clauses max : assez pour le rôle, sans coller tout le paragraphe.
    kept = clauses[:2]
    out = ". ".join(c.rstrip(".") for c in kept).rstrip(".") + "."
    out = re.sub(r"\s+", " ", out).strip()
    if len(out) > 180:
        out = out[:177].rsplit(" ", 1)[0].rstrip(".,;") + "."
    return out or None


def detect_series_brand(title):
    up = clean_spaces(title).upper()
    if up in SERIES_BRAND:
        return SERIES_BRAND[up]
    for key, brand in SERIES_BRAND.items():
        if up.startswith(key):
            return brand
    return ""


def looks_like_product_title(line):
    raw = clean_spaces(line)
    line = re.sub(r"^\d{1,4}(?=[A-ZÀ-Ÿ])", "", raw).strip()
    reject = None
    if is_slogan_line(line):
        reject = "slogan"
    elif is_table_section_header(line):
        reject = "table_section"
    elif line.upper() in TITLE_STOPWORDS:
        reject = "stopword"
    elif line.startswith("-") or line.startswith("–") or line.startswith("—"):
        reject = "dash"
    elif CODE_RE.search(line) or PRICE_RE.search(line):
        reject = "code_or_price"
    elif HEADER_ROW_RE.search(line) or is_skip_line(line):
        reject = "skip_or_header"
    elif is_warranty_or_meta_line(line):
        reject = "warranty"
    elif is_category_header(line):
        reject = "category"
    elif is_junk_title(line):
        reject = "junk"
    else:
        up = line.upper()
        if any(k in up for k in ("PRODUITS DISPONIBLES", "6 000", "6000 PRODUITS", "EN AGENCE")):
            reject = "promo"
        else:
            letters = [c for c in line if c.isalpha()]
            if not letters:
                reject = "no_letters"
            else:
                upper_ratio = sum(c.isupper() for c in letters) / len(letters)
                short_name = (
                    3 <= len(line) <= 5
                    and line.isupper()
                    and len(letters) >= 3
                    and not any(c.isdigit() for c in line)
                )
                if short_name:
                    reject = None
                elif not (6 <= len(line) <= 120):
                    reject = f"len:{len(line)}"
                elif upper_ratio < 0.45:
                    reject = "lowercase"
    if reject:
        return False
    return True


def non_category_texts(fitz_page):
    """Textes à ne pas confondre avec une rubrique.

    Renvoie (titres, marge) : les titres produit repérés à la typographie, et
    les onglets pivotés de la marge, que pdfium recolle parfois en une seule
    ligne (« EQUIPEMENT DE CHAUFFAGE GÉNIE CLIMATIQUE »).
    """
    titles, margin = set(), set()
    for block in fitz_page.get_text("dict").get("blocks") or []:
        if block.get("type") != 0:
            continue
        for line in block.get("lines") or []:
            spans = line.get("spans", [])
            text = clean_spaces("".join(s.get("text") or "" for s in spans))
            if not text:
                continue
            if tuple(line.get("dir") or (1, 0)) != (1.0, 0.0):
                margin.add(text.upper())
            elif line_role(spans) == "title":
                titles.add(text.upper())
    return titles, margin


def is_non_category(line, titles=(), margin=()):
    """Vrai si la ligne est un titre produit ou reprend un onglet de marge."""
    up = clean_spaces(line).upper()
    if up in titles:
        return True
    return any(tab in up for tab in margin)


def detect_category(page_text, pending="", titles=(), margin=()):
    lines = [clean_spaces(x) for x in page_text.splitlines() if clean_spaces(x)]
    for line in lines[:6]:
        if is_skip_line(line) or CODE_RE.search(line) or is_slogan_line(line):
            continue
        if line[:1].islower() or " - " in line:
            continue
        line2 = re.sub(r"^\d{1,4}(?=[A-ZÀ-Ÿ ])", "", line).strip()
        if is_non_category(line2, titles, margin):
            continue
        if 4 <= len(line2) <= 28 and line2.isupper() and "Ø" not in line2:
            if any(
                k in line2
                for k in ("TUBE", "RACCORD", "ROBINET", "COLLECTEUR", "PLOMBERIE", "VANNE")
            ):
                return line2
    candidates = []
    for line in lines[:12]:
        if is_skip_line(line) or CODE_RE.search(line) or is_slogan_line(line):
            continue
        if line[:1].islower() or " - " in line:
            continue
        # Une rubrique n'est pas une phrase : sans ce garde-fou la description
        # courte d'un produit finit en catégorie.
        if looks_like_description_line(line):
            continue
        line2 = re.sub(r"^\d{1,4}(?=[A-ZÀ-Ÿ ])", "", line).strip()
        up = line2.upper()
        if is_non_category(line2, titles, margin):
            continue
        if any(
            k in up
            for k in (
                "RECEVEUR", "RACCORD", "TUBE", "ROBINET", "CHAUFFE", "RADIATEUR",
                "PAROI", "WC", "LAVABO", "DOUCHE", "MITIGEUR", "POMPE", "CLIM",
                "VENTIL", "CARRELAGE", "OUTILLAGE", "LAITON", "CUIVRE",
                "MULTICOUCHE", "ABATTANT",
            )
        ) and 8 <= len(line2) <= 80:
            candidates.append(line2)
    if candidates:
        scored = []
        for s in candidates:
            has_diam = "Ø" in s or bool(re.search(r"\d+\s*MM", s, re.I))
            scored.append((s.isupper(), not has_diam, len(s) < 55, -len(s), s))
        return max(scored)[-1]
    return pending


def merge_wrapped_lines(lines):
    out = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        if not CODE_RE.search(line) and i + 1 < n:
            nxt = lines[i + 1]
            join = False
            if WRAP_END_RE.search(line):
                join = True
            elif looks_like_row_prefix(line) and CODE_RE.search(nxt):
                join = True
            elif re.match(r"^(cm|mm)\b", nxt, re.I) and re.search(r"\d", line):
                join = True
            if join:
                combined = clean_spaces(line + " " + nxt)
                i += 1
                if not CODE_RE.search(combined) and i + 1 < n and CODE_RE.search(lines[i + 1]):
                    combined = clean_spaces(combined + " " + lines[i + 1])
                    i += 1
                out.append(combined)
                i += 1
                continue
        out.append(line)
        i += 1
    return out


def is_supplier_ref(tok):
    if re.fullmatch(r"\d{6,}", tok):
        return True
    if not re.fullmatch(r"[A-Z0-9._-]{6,}", tok, re.I):
        return False
    if not re.search(r"[A-Za-z]", tok):
        return False
    if not re.search(r"\d", tok):
        return False
    if re.fullmatch(r"\d{1,2}(?:[-/]\d{1,2})?(?:[xX×]\d{1,2}){1,2}(?:-\d{1,2})?", tok):
        return False
    return True


def parse_product_row(line, current_vendu=""):
    line = clean_spaces(line)
    if HEADER_ROW_RE.search(line) and not CODE_RE.search(line):
        return None
    codes = CODE_RE.findall(line)
    if not codes:
        return None
    code = codes[-1]
    if len(set(code)) == 1:
        return None

    before, after = line.rsplit(code, 1)
    before = before.strip(" -•●")
    after = after.strip()

    price = ""
    if PRICE_COURS_RE.search(after) or PRICE_COURS_RE.search(line):
        price = "Prix à cours"
    else:
        pm = PRICE_RE.search(after) or PRICE_RE.search(line)
        if pm:
            price = normalize_price(pm.group(1))

    tokens = before.split()
    variant_letter = ""
    if tokens and re.fullmatch(r"[abcABC●•]", tokens[0]):
        variant_letter = tokens[0].upper().replace("●", "").replace("•", "")
        if variant_letter in "ABCD":
            tokens = tokens[1:]
        else:
            variant_letter = ""

    reference = ""
    ref_four = ""
    vendu = current_vendu
    diametre = ""
    designation_extra = ""

    if tokens:
        idx = len(tokens) - 1
        while idx >= 0:
            tok = tokens[idx]
            if re.fullmatch(r"\d{3,5}", tok):
                reference = tok
                tokens = tokens[:idx]
                break
            if is_supplier_ref(tok):
                ref_four = tok
                tokens = tokens[:idx]
                idx -= 1
                continue
            break

        while tokens and is_supplier_ref(tokens[-1]):
            ref_four = tokens[-1]
            tokens = tokens[:-1]
        while tokens and re.fullmatch(r"\d+(?:\.\d+){1,4}", tokens[-1]):
            ref_four = tokens[-1]
            tokens = tokens[:-1]
        if tokens and re.fullmatch(r"\d{1,3},\d{1,2}", tokens[-1]):
            tokens = tokens[:-1]
        if (
            len(tokens) >= 2
            and re.fullmatch(r"(1|2|5|10|20|25|50)", tokens[-1])
            and re.search(r"[A-Za-zÀ-Ÿ]", " ".join(tokens[:-1]))
        ):
            vendu = tokens[-1]
            tokens = tokens[:-1]

        rest = " ".join(tokens).strip()
        if rest:
            compact = rest.replace(" ", "")
            if DIAM_RE.match(compact) or DIAM_RE.match(rest):
                diametre = compact.replace("×", "X").replace("x", "X")
            else:
                designation_extra = rest
                sm = SIZE_RE.search(rest)
                if sm:
                    diametre = clean_spaces(sm.group(1))
                    diametre = re.sub(r"\s*[xX×]\s*", " x ", diametre)

    if not price and not designation_extra and not diametre and not reference:
        return None

    return {
        "Réf.Pro": reference,
        "Réf.Four": ref_four,
        "Diamètre": diametre,
        "Vendu par": vendu,
        "extra": designation_extra,
        "Code": code,
        "Prix HT": price,
        "letter": variant_letter,
    }


TITLE_CONTINUE_RE = re.compile(
    r'^[\-–—•●"]|^(AVEC|GARDE|SORTIE|POUR|ANTI)\b',
    re.I,
)
DIM_LINE_RE = re.compile(
    r"^Ø?\s*\d+[.,]?\d*\s*(mm)?(\s*[–\-xX×]\s*)?(L\.\s*\d+\s*m)?$",
    re.I,
)


def is_title_continuation(line):
    s = clean_spaces(line)
    if not s or CODE_RE.search(s) or PRICE_RE.search(s) or is_slogan_line(s):
        return False
    if looks_like_description_line(s):
        return False
    if re.search(r"pression|raccordement|filetage|temp[eé]rature|garantie", s, re.I):
        return False
    if s.upper() in TITLE_CONTINUE_WORDS:
        return True
    if len(s) > 48:
        return False
    letters = [c for c in s if c.isalpha()]
    if letters and sum(c.islower() for c in letters) / len(letters) > 0.3:
        return False
    if DIM_LINE_RE.match(s.replace(" ", "")) or re.match(r"^Ø\s*\d", s, re.I):
        return False
    if is_junk_title(s.lstrip("-–—•●\" ")):
        return False
    return bool(TITLE_CONTINUE_RE.match(s))


JUNK_TITLE_RE = re.compile(
    r"^(pression|temp[eé]rature|service\s*:|classe\s+\d|pd\s*=|d[eé]signation|"
    r"eco-contribution|le \+|bao$|article top|barri[eè]re anti)",
    re.I,
)


def is_junk_title(line):
    """Specs / fins de phrase (ex. ': 95°C') ≠ titre de fiche."""
    s = clean_spaces(line)
    if not s:
        return True
    if is_slogan_line(s):
        return True
    if re.fullmatch(r"\d+\s*(MM|CM)", s, re.I):
        return True
    if s.startswith(":") or s.startswith("–") or s.startswith("—"):
        return True
    if re.match(r"^[●•▪]", s):
        return True
    if JUNK_TITLE_RE.match(s):
        return True
    if re.search(r"accidentelle|max de service|pression maximale|temp[eé]rature max", s, re.I):
        return True
    if re.search(r"diamètre\s+de\s+raccord|des\s+circuits\s+eurokonus", s, re.I):
        return True
    if re.search(r"\b(plancher chauffant|rafraîchissant|eurokonus|entraxe)\b", s, re.I):
        return True
    if re.search(r"\d+\s+circuits\s+\d+,\d{2}\s+\d{6,}", s, re.I):
        return True
    if re.match(r"^\d+\s+circuits\b", s, re.I):
        return True
    if re.fullmatch(r"GARANTIE\.?", s, re.I):
        return True
    letters = [c for c in s if c.isalpha()]
    if len(letters) < 3:
        return True
    return False


def peel_technical(title):
    """Garde le début de phrase ; coupe seulement une queue trop longue."""
    t = clean_spaces(title).strip(" -–—:")
    t = t.replace('"', "").strip()
    diameter = ""
    notes = []

    dm = re.search(
        r"(Ø\s*\d+[.,]?\d*\s*(?:MM|CM)?(?:\s*[–\-xX×]\s*\d+[.,]?\d*\s*(?:MM|CM)?)?)",
        t,
        re.I,
    )
    lm = re.search(r"(L\.\s*\d+\s*m)", t, re.I)
    if dm:
        diameter = clean_spaces(dm.group(1))
    if lm:
        diameter = clean_spaces(f"{diameter} {lm.group(1)}".strip())

    if len(t) > MAX_NAME_LEN:
        cut = t[:MAX_NAME_LEN].rsplit(" ", 1)[0]
        tail = clean_spaces(t[len(cut) :])
        if tail:
            notes.append(tail)
        t = cut

    note = " · ".join(n.strip(" -–") for n in notes if n and n.strip(" -–"))
    return t, diameter, note


def clean_product_title(title):
    """Titre de fiche produit (partagé par toutes les variantes)."""
    t = clean_spaces(title or "")
    t = re.sub(r"\s+[Bb]\d{4,6}\s*$", "", t)
    t = re.sub(r"\s*\(paire\)\s*$", "", t, flags=re.I)
    name, _diam, _note = peel_technical(t)
    return name or t


def build_designation(category, title, extra):
    """Nom clair : début de phrase uniquement. extra (Ø, L., réf) ne va pas dans le nom."""
    tit = clean_spaces(title)
    cat = clean_spaces(category)
    base = tit or cat
    name, _diam, _note = peel_technical(base)
    return name or base


def looks_like_dimension_line(line):
    s = clean_spaces(line)
    if re.match(r"^Ø\s*\d", s, re.I):
        return True
    if re.search(r"L\.\s*\d+\s*m", s, re.I) and len(s) < 40 and not re.search(r"[A-Za-zÀ-Ÿ]{5,}", s.replace("mm", "")):
        return True
    return False


def extract_page_text(page):
    tp = page.get_textpage()
    return tp.get_text_bounded() if tp else ""


def find_variant_letter(fitz_page, code, page_h):
    """Repère le marqueur A/B/C à gauche d'une ligne de code dans le PDF."""
    code_bbox = None
    for b in fitz_page.get_text("dict")["blocks"]:
        if b.get("type") != 0:
            continue
        for line in b["lines"]:
            txt = "".join(s["text"] for s in line["spans"]).strip()
            if code in txt:
                code_bbox = line["bbox"]
                break
        if code_bbox:
            break
    if not code_bbox:
        return ""
    code_y_mid = (code_bbox[1] + code_bbox[3]) / 2
    code_x = code_bbox[0]
    best = ""
    best_dist = 999.0
    for b in fitz_page.get_text("dict")["blocks"]:
        if b.get("type") != 0:
            continue
        for line in b["lines"]:
            for span in line["spans"]:
                t = span["text"].strip()
                if not re.fullmatch(r"[A-Da-d]", t):
                    continue
                sb = span["bbox"]
                sy = (sb[1] + sb[3]) / 2
                if abs(sy - code_y_mid) > 10 or sb[2] > code_x:
                    continue
                dist = code_x - sb[2]
                if dist < best_dist:
                    best_dist = dist
                    best = t.upper()
    return best


def collect_block_photo_row(photos, y0, y1, col_left, page_w, sku_y_min):
    """Photos côte à côte cohérentes dans le bloc article (au-dessus des lignes SKU)."""
    if not photos or y0 is None:
        return []
    col_photos = [
        p
        for p in photos
        if (p["x"] < page_w / 2) == col_left and (y0 - 12) <= p["y"] < (y1 - 6)
    ]
    # Pas de repli sur toute la colonne : cela volait les photos de l'article
    # suivant. Mieux vaut aucune photo qu'une photo d'un autre produit.
    if sku_y_min is not None:
        above = [p for p in col_photos if p["y"] < sku_y_min - 8]
        if above:
            col_photos = above
    if not col_photos:
        return []

    # Les photos d'une même rangée sont alignées par le bas, sur la ligne des
    # étiquettes A/B/C : une pièce plate a le même bas mais un haut bien plus
    # bas qu'une pièce haute. Regrouper par le haut la laissait de côté. On
    # garde la rangée la plus fournie : un petit visuel isolé (flash
    # « NOUVEAU ») ne doit pas servir de référence.
    by_base = []
    for photo in sorted(col_photos, key=lambda p: p["y1"]):
        if by_base and abs(photo["y1"] - by_base[-1][0]["y1"]) <= 14:
            by_base[-1].append(photo)
        else:
            by_base.append([photo])
    band = max(by_base, key=lambda r: (len(r), r[0]["y1"]))
    band.sort(key=lambda p: p["x"])
    if len(band) == 1:
        return band

    runs = [[band[0]]]
    for photo in band[1:]:
        prev = runs[-1][-1]
        if photo["x"] - prev["x"] <= 85:
            runs[-1].append(photo)
        else:
            runs.append([photo])
    best = max(runs, key=lambda r: (len(r), -r[0]["y1"]))
    return best


def photo_letter_index(letter):
    if not letter:
        return None
    letter = letter.upper()
    if letter in "ABCD":
        return ord(letter) - ord("A")
    return None


def assign_block_photos(items, photos, fitz_page, tp, ph, pw, left_ys, right_ys):
    """Associe les photos A/B/C aux variantes d'un même bloc article."""
    from collections import defaultdict

    groups = defaultdict(list)
    for item in items:
        art_y = item.get("_art_y0")
        art_col = item.get("_art_col")
        ty, tx = find_title_y(tp, (item.get("Désignation") or "")[:28], ph)
        if art_col is None:
            art_col = 0 if (tx or 0) < pw / 2 else 1
        if art_y is None:
            art_y = ty
        # Un même titre (« COUPE-TUBE MÉTAL ») peut apparaître 4 fois sur la page :
        # on groupe par bloc (colonne + Y du titre), pas par nom.
        groups[(round(float(art_y or 0), 0), int(art_col))].append((item, ty, tx))

    for (_art_y, col), rows in groups.items():
        if not rows:
            continue
        tx = 40.0 if int(col) == 0 else pw * 0.75
        col_left = int(col) == 0
        code_ys = []
        for item, _ty, _tx in rows:
            cy, _ = find_title_y(tp, item["Code"], ph)
            if cy is not None:
                code_ys.append(cy)
        if not code_ys:
            continue
        sku_y_min = min(code_ys)
        y0, y1 = band_for(sku_y_min, tx, left_ys, right_ys, pw)
        if y0 is None:
            y0 = max(0.0, sku_y_min - 140)
        if y1 is None:
            y1 = min(ph, max(code_ys) + 50)

        block_photos = collect_block_photo_row(photos, y0, y1, col_left, pw, sku_y_min)
        hero = block_photos[0]["file"] if block_photos else ""
        if len(block_photos) < 2:
            for item, _ty, _tx in rows:
                fname = hero or pick_photo_in_block(photos, y0, y1, tx, sku_y_min, pw)
                if fname:
                    item["Image"] = f"{IMAGE_REL}/{fname}"
                item["_hero"] = fname or item.get("_hero") or ""
            continue

        for item, _ty, _tx in rows:
            letter = item.get("_photo_letter") or find_variant_letter(fitz_page, item["Code"], ph)
            idx = photo_letter_index(letter)
            if idx is None or idx >= len(block_photos):
                # Sans lettre A/B/C : photo principale du bloc (A), jamais la
                # molette / accessoire collé au tableau.
                fname = block_photos[0]["file"]
                item["Image"] = f"{IMAGE_REL}/{fname}"
                item["_hero"] = hero
                continue
            item["Image"] = f"{IMAGE_REL}/{block_photos[idx]['file']}"
            item["_hero"] = hero


def pad_white(pil, frac=PAD_FRAC):
    """Marge blanche pour éviter un recadrage trop serré (effet zoom)."""
    w, h = pil.size
    mx = max(8, int(w * frac))
    my = max(8, int(h * frac))
    canvas = Image.new("RGB", (w + 2 * mx, h + 2 * my), (255, 255, 255))
    canvas.paste(pil, (mx, my))
    return canvas


def person_photo_features(pil):
    thumb = pil.resize((64, 64))
    pix = list(thumb.getdata())
    n = len(pix)
    white = skin = 0
    bins = Counter()
    for r, g, b in pix:
        if r > 230 and g > 230 and b > 230:
            white += 1
            continue
        if r > 95 and g > 40 and b > 20 and r > g > b and (r - g) > 15 and (r - b) > 20 and r < 230:
            skin += 1
        bins[(r // 32, g // 32, b // 32)] += 1
    white_r = white / n
    skin_r = skin / n
    total = sum(bins.values()) or 1
    entropy = 0.0
    for c in bins.values():
        p = c / total
        entropy -= p * math.log2(p)
    return white_r, skin_r, entropy


def looks_like_person(pil):
    """Photos de personnes / lifestyle — pas les pièces (laiton ≠ visage)."""
    w, h = pil.size
    white, skin, entropy = person_photo_features(pil)
    feats = {"white": round(white, 3), "skin": round(skin, 3), "entropy": round(entropy, 2)}
    # Vignettes catalogue (emballage coloré type tuyau) ≠ photos de personnes
    if max(w, h) <= 165:
        return False, feats
    # Visages : peau modérée + couleurs variées. Les pièces laiton ont skin > 0.5.
    if 0.035 <= skin <= 0.42 and entropy >= 5.0 and white < 0.50:
        return True, {"white": round(white, 3), "skin": round(skin, 3), "entropy": round(entropy, 2)}
    # Portraits / scènes large (équipe, showroom)
    if w >= 300 and h >= 220 and 0.04 <= skin <= 0.40 and entropy >= 4.5 and white < 0.50:
        return True, {"white": round(white, 3), "skin": round(skin, 3), "entropy": round(entropy, 2), "large": True}
    return False, {"white": round(white, 3), "skin": round(skin, 3), "entropy": round(entropy, 2)}


def save_hq_image(pil, dest):
    """PNG à taille catalogue web — pas d'agrandissement, réduction si trop grand."""
    w, h = pil.size
    longest = max(w, h)
    if longest > MAX_IMAGE_PX:
        scale = MAX_IMAGE_PX / longest
        pil = pil.resize((max(1, int(w * scale)), max(1, int(h * scale))), Image.Resampling.LANCZOS)
    pil.save(dest, "PNG", optimize=True)


def looks_like_logo_photo(dw, dh, px_w, px_h):
    """Logos / bandeaux marque : très larges et bas, pas une photo de pièce."""
    if dh <= 0:
        return True
    aspect = dw / dh
    px_aspect = max(px_w, px_h) / max(1, min(px_w, px_h))
    if aspect >= 2.4 and dh < 26:
        return True
    # Wordmark de marque (Virax, Rothenberger, Knipex…) : plus étroit qu'une pièce.
    if dw < 50 and dh < 24 and aspect >= 1.7:
        return True
    if px_aspect >= 3.0 and min(px_w, px_h) < 80:
        return True
    if dh < 14:
        return True
    return False


def extract_product_photos(page, page_idx, image_dir, hash_to_name, skip_stats):
    """Extrait les photos natives des pièces (pas de rendu ×3, pas de personnes)."""
    pw, ph = page.get_size()
    # get_size() donne la CropBox, mais get_bounds() renvoie des coordonnées
    # MediaBox brutes. fitz, lui, place l'origine du texte sur la CropBox. Sans
    # cette conversion les photos sont décalées de la marge (~30 pt) par rapport
    # au texte, et se retrouvent rattachées à l'article voisin.
    try:
        crop_x0, _cy0, _cx1, crop_y1 = page.get_cropbox()
    except Exception:
        crop_x0, crop_y1 = 0.0, ph
    photos = []
    img_i = 0

    for obj in page.get_objects():
        if type(obj).__name__ != "PdfImage":
            continue
        try:
            px_w, px_h = obj.get_px_size()
            bounds = obj.get_bounds()
        except Exception:
            continue
        x0, y0, x1, y1 = bounds
        dw, dh = abs(x1 - x0), abs(y1 - y0)
        if min(px_w, px_h) < 48 or max(px_w, px_h) < 90:
            if max(px_w, px_h) < 70 and max(dw, dh) < 40:
                continue
        if looks_like_logo_photo(dw, dh, px_w, px_h):
            continue
        aspect = max(px_w, px_h) / max(1, min(px_w, px_h))
        max_aspect = 5.2 if max(px_w, px_h) >= 100 else 3.2
        if aspect > max_aspect:
            continue
        y_top = crop_y1 - max(y0, y1)
        if y_top < ph * 0.04 or y_top > ph * 0.92:
            continue

        try:
            pil = obj.get_bitmap().to_pil().convert("RGB")
        except Exception:
            continue
        if min(pil.width, pil.height) < 48 or max(pil.width, pil.height) < 90:
            if max(pil.width, pil.height) < 70 and max(dw, dh) < 40:
                continue

        is_person, feats = looks_like_person(pil)
        if is_person:
            skip_stats["people"] += 1
            continue

        pil = pad_white(pil)

        raw = pil.tobytes()
        digest = hashlib.md5(raw).hexdigest()
        if digest in hash_to_name:
            fname = hash_to_name[digest]
        else:
            fname = f"p{page_idx + 1:04d}_{img_i:02d}.{IMAGE_EXT}"
            dest = os.path.join(image_dir, fname)
            save_hq_image(pil, dest)
            hash_to_name[digest] = fname
            img_i += 1
            skip_stats["kept"] += 1

        photos.append(
            {
                "x": (x0 + x1) / 2 - crop_x0,
                "y": y_top,
                "y1": y_top + abs(y1 - y0),
                "w": abs(x1 - x0),
                "h": abs(y1 - y0),
                "file": fname,
            }
        )

    photos.sort(key=lambda p: (p["x"] > pw / 2, p["y"]))
    return photos


def find_text_pos_fitz(fitz_page, needle):
    """Position (y, x) d'un texte dans le PDF (origine haut, comme les traits verts)."""
    if not needle or not fitz_page:
        return None, None
    for block in fitz_page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            text = "".join(span.get("text", "") for span in line.get("spans", []))
            if needle in text:
                bbox = line.get("bbox", (0, 0, 0, 0))
                return float(bbox[1]), float(bbox[0])
    return None, None


def find_title_y(tp, title, page_h):
    if not title or not tp:
        return None, None
    needle = title[:24]
    try:
        text = tp.get_text_range().replace("\r", " ").replace("\n", " ")
    except Exception:
        return None, None
    pos = text.find(needle)
    if pos < 0:
        pos = text.upper().find(needle.upper())
    if pos < 0:
        return None, None
    try:
        box = tp.get_charbox(pos)
        return page_h - box[3], box[0]
    except Exception:
        return None, None


def nearest_photo(photos, title_y, title_x_hint, page_w):
    if not photos:
        return ""
    col_left = title_x_hint is None or title_x_hint < page_w / 2
    same = [p for p in photos if (p["x"] < page_w / 2) == col_left]
    pool = same or photos
    if title_y is None:
        return pool[0]["file"]
    best = min(
        pool,
        key=lambda p: abs(p["y"] - title_y) + abs(p["x"] - (title_x_hint or p["x"])) * 0.15,
    )
    return best["file"]


def article_separator_ys(fitz_page, page_w, page_h):
    """Traits rouges ou verts = limites d'articles dans le catalogue CEDEO."""
    left, right = [0.0], [0.0]
    try:
        drawings = fitz_page.get_drawings()
    except Exception:
        return [0.0, page_h], [0.0, page_h]
    for d in drawings:
        color = d.get("color") or d.get("fill")
        if not color or len(color) < 3:
            continue
        r, g, b = color[:3]
        is_red = r >= 0.65 and g < 0.35 and b < 0.35
        is_green = r < 0.15 and 0.45 < g < 0.75 and b < 0.45
        if not (is_red or is_green):
            continue
        rect = d.get("rect")
        if rect is None or rect.height > 4 or rect.width < 80:
            continue
        y = float(rect.y0)
        cx = (float(rect.x0) + float(rect.x1)) / 2
        if cx < page_w / 2:
            left.append(y)
        else:
            right.append(y)
    left.append(page_h)
    right.append(page_h)
    def uniq(xs):
        out = []
        for y in sorted(xs):
            if not out or abs(y - out[-1]) > 8:
                out.append(y)
        return out
    return uniq(left), uniq(right)


def red_separator_ys(fitz_page, page_w, page_h):
    """Alias conservé pour compatibilité interne."""
    return article_separator_ys(fitz_page, page_w, page_h)


def band_for(y, x, left_ys, right_ys, page_w):
    if y is None:
        return None, None
    ys = left_ys if (x is None or x < page_w / 2) else right_ys
    y0 = ys[0]
    y1 = ys[-1]
    for i in range(len(ys) - 1):
        if ys[i] - 4 <= y < ys[i + 1]:
            return ys[i], ys[i + 1]
    return y0, y1


def pick_photo_in_block(photos, y0, y1, title_x, sku_y, page_w):
    """Photo du bloc : à droite / au-dessus des désignations (bas de photo ≈ tableau)."""
    if not photos or y0 is None:
        return ""
    col_left = title_x is None or title_x < page_w / 2
    cand = [
        p
        for p in photos
        if (y0 - 12) <= p["y"] < (y1 - 6) and (p["x"] < page_w / 2) == col_left
    ]
    if not cand:
        cand = [p for p in photos if (y0 - 12) <= p["y"] < (y1 - 6)]
    if not cand:
        return ""
    if sku_y is not None:
        above = [p for p in cand if p["y"] <= sku_y + 8]
        pool = above or cand
        # Plus grande photo du bloc (pièce), pas celle collée au tableau
        # (molette, joint, pictogramme).
        return max(
            pool,
            key=lambda p: (
                (p.get("w") or 40) * (p.get("h") or 40),
                -(p.get("y") or 0),
            ),
        )["file"]
    return max(cand, key=lambda p: (p.get("w") or 0) * (p.get("h") or 0))["file"]


def looks_like_running_header(text, y):
    """Bandeau de page (LAVE-MAINS en haut) ≠ nom d'article au-dessus de la photo."""
    if y is None or y > 30:
        return False
    s = clean_spaces(text)
    if not s:
        return True
    if is_category_header(s):
        return True
    up = s.upper()
    if up in {"LAVE-MAINS", "SANITAIRE", "PLOMBERIE", "COIN TOILETTES", "ACCESSIBILITÉ", "COLLECTIVITÉ"}:
        return True
    if len(s) <= 24 and s.isupper() and not re.search(r"\d", s) and "MEUBLE" not in up:
        return True
    return False


def column_starts_new_article(items):
    """False = suite de tableau / description (article de la page précédente)."""
    for y, _x0, text, role in items:
        if is_skip_line(text) or is_warranty_or_meta_line(text) or is_slogan_line(text):
            continue
        if is_layout_chrome(text) or looks_like_row_prefix(text):
            continue
        if looks_like_running_header(text, y):
            continue
        if role == "title":
            return True
        if CODE_RE.search(text) or PRICE_RE.search(text):
            return False
        if is_table_section_header(text):
            return False
        if role == "desc":
            return False
    return False


def collect_page_articles(fitz_page, page_w, page_h, category="", carry=None, page_no=0):
    """Articles : colonne gauche puis droite. Suite page précédente si pas de nouveau titre."""
    left_ys, right_ys = article_separator_ys(fitz_page, page_w, page_h)
    articles = []
    col_items = extract_column_items(fitz_page, page_w)
    starts_new = column_starts_new_article(col_items["L"])
    if carry and not starts_new:
        carried = dict(carry)
        carried["col"] = "L"
        carried["y0"] = 0.0
        carried["y1"] = band_end_for(0.0, left_ys) or page_h
        carried["sections"] = []
        carried["variant_group"] = ""
        carried["variant_sub"] = ""
        articles.append(carried)
    # Y des lignes de tableau (celles portant un code) : leurs autres cellules
    # (Réf. Four., désignation) ne sont ni des titres ni des descriptions.
    row_ys = {
        col: [y for y, _x, t, _r in col_items[col] if CODE_RE.search(t)]
        for col in ("L", "R")
    }

    def on_table_row(col, y):
        return any(abs(y - ry) <= 4 for ry in row_ys[col])

    for col in ("L", "R"):
        seps = left_ys if col == "L" else right_ys
        current = None
        if col == "L" and articles and carry and articles[0].get("y0") == 0.0:
            current = articles[0]
        for y, x0, text, role in col_items[col]:
            if is_skip_line(text) or is_warranty_or_meta_line(text) or is_slogan_line(text):
                continue
            if is_layout_chrome(text) or looks_like_row_prefix(text):
                continue
            if looks_like_running_header(text, y):
                if not category:
                    category = clean_spaces(text)
                continue
            if current and current.get("y1") is not None and y >= current["y1"] - 4:
                current = None
            # La typographie primant sur tout le reste, le titre est traité avant
            # les tests de contenu : aucun d'eux ne doit pouvoir l'avaler.
            if role == "title":
                if (
                    current
                    and current.get("_title_y") is not None
                    and y - current["_title_y"] <= 14
                ):
                    # Titre sur deux lignes : elles se suivent à ~10 pt d'écart.
                    current["title"] = clean_spaces(current["title"] + " " + text)
                    current["_title_y"] = y
                    continue
                title = re.sub(r"^\d{1,4}(?=[A-ZÀ-Ÿ])", "", text).strip()
                current = {
                    "title": title,
                    "description": "",
                    "brand": detect_series_brand(title) or detect_brand(title) or "",
                    "category": category,
                    "y0": y,
                    "y1": band_end_for(y, seps) or page_h,
                    "col": col,
                    "x": x0,
                    "variant_group": "",
                    "variant_sub": "",
                    "sections": [],
                    "note": "",
                    "_title_y": y,
                }
                articles.append(current)
                continue
            if CODE_RE.search(text) or PRICE_RE.search(text):
                continue
            if is_category_header(text, category):
                category = clean_spaces(text)
                if current:
                    current["category"] = category
                continue
            bm = BULLET_NOTE_RE.match(text)
            if bm and current and not is_table_section_header(text):
                note = clean_spaces(bm.group(1))
                if note and not CODE_RE.search(note):
                    current["note"] = note.upper()
                continue
            if is_table_section_header(text):
                if current:
                    apply_table_section(current, text)
                    current["sections"].append(
                        {
                            "y": y,
                            "group": current.get("variant_group") or "",
                            "sub": current.get("variant_sub") or "",
                        }
                    )
                continue
            if on_table_row(col, y):
                continue
            if role == "desc":
                # Le corps 6.50 sert aussi aux intitulés de rubrique sous le titre
                # (« CHAUFFAGE »). Une vraie description est une phrase.
                # Exception : dernière ligne tout en chiffres/majuscules
                # (« 2300 / S 2400 / C 2300 G ») si on a déjà commencé le paragraphe.
                if re.search(r"\d\s*x\s*\d+\s*V|\d+\s*Hz|1x230", text, re.I):
                    continue
                has_lower = any(c.islower() for c in text)
                if not has_lower and not (current and current.get("description")):
                    if current and not current.get("category"):
                        current["category"] = clean_spaces(text)
                    continue
                if current:
                    current["description"] = clean_spaces(
                        (current.get("description") or "") + " " + text
                    )
                continue
            if current and looks_like_dimension_line(text):
                current["last_diam"] = clean_spaces(text)
    next_carry = None
    right_arts = [a for a in articles if a.get("col") == "R"]
    candidate = right_arts[-1] if right_arts else (articles[-1] if articles else None)
    if candidate and (candidate.get("y1") or 0) >= page_h - 8:
        next_carry = candidate
    return articles, category, next_carry


def find_article_for_code(articles, y, x, page_w):
    col = "L" if (x or 0) < page_w / 2 else "R"
    y = y if y is not None else 0
    hits = [
        a
        for a in articles
        if a.get("col") == col and (a.get("y0") or 0) - 8 <= y < (a.get("y1") or 0) + 6
    ]
    if hits:
        return max(hits, key=lambda a: a.get("y0") or 0)
    prev = [a for a in articles if a.get("col") == col and (a.get("y0") or 0) <= y + 8]
    return prev[-1] if prev else None


def section_for_article(art, y):
    if not art:
        return None
    y = y if y is not None else 0
    secs = [s for s in (art.get("sections") or []) if s.get("y", 0) <= y + 4]
    return secs[-1] if secs else None


def process_pdf():
    if not os.path.isfile(PDF_PATH):
        print(f"PDF introuvable: {PDF_PATH}", file=sys.stderr)
        sys.exit(1)

    os.makedirs(IMAGE_DIR, exist_ok=True)
    pdf = pdfium.PdfDocument(PDF_PATH)
    fitz_doc = fitz.open(PDF_PATH)
    products = []
    seen = set()
    hash_to_name = {}
    skip_stats = {"people": 0, "kept": 0, "logged": 0, "size_logged": 0}
    pending = {"category": ""}
    article_carry = None

    page_from = int(os.environ.get("EXTRACT_PAGE_FROM") or "1")
    page_to = int(os.environ.get("EXTRACT_PAGE_TO") or str(len(pdf)))
    only_pages = set()
    raw_pages = os.environ.get("EXTRACT_PAGES") or ""
    if raw_pages.strip():
        for part in raw_pages.replace(" ", "").split(","):
            if not part:
                continue
            if "-" in part:
                a, b = part.split("-", 1)
                only_pages.update(range(int(a), int(b) + 1))
            else:
                only_pages.add(int(part))

    for idx in range(len(pdf)):
        page_no = idx + 1
        if only_pages and page_no not in only_pages:
            continue
        if page_no < page_from or page_no > page_to:
            continue
        page = pdf[idx]
        text = extract_page_text(page)
        if not text.strip():
            continue
        pw, ph = page.get_size()
        tp = page.get_textpage()
        page_start = len(products)
        fitz_page = fitz_doc[idx]
        fw, fh = fitz_page.rect.width, fitz_page.rect.height

        cat_titles, cat_margin = non_category_texts(fitz_page)
        cat = detect_category(text, pending["category"], cat_titles, cat_margin)
        if cat:
            pending["category"] = cat

        if os.environ.get("EXTRACT_TEXT_ONLY"):
            photos = []
        else:
            photos = extract_product_photos(page, idx, IMAGE_DIR, hash_to_name, skip_stats)

        articles, pending["category"], article_carry = collect_page_articles(
            fitz_page, fw, fh, pending["category"], article_carry, page_no
        )
        for art in articles:
            art["image"] = nearest_photo(photos, art.get("y0"), art.get("x"), pw)

        raw_lines = [clean_spaces(x) for x in text.splitlines() if clean_spaces(x)]
        lines = merge_wrapped_lines(raw_lines)
        last_prefix = ""

        for line in lines:
            if looks_like_row_prefix(line):
                last_prefix = line
                continue
            parsed = parse_product_row(line, "")
            if not parsed:
                continue
            code = parsed["Code"]
            if code in seen:
                continue

            cy, cx = find_text_pos_fitz(fitz_page, code)
            if cy is None or cx is None:
                if len(articles) == 1:
                    art = articles[0]
                    cy = articles[0].get("y0")
                    cx = articles[0].get("x")
                else:
                    continue
            else:
                art = find_article_for_code(articles, cy, cx, fw)
            if not art or not art.get("title") or is_slogan_line(art.get("title") or ""):
                continue
            seen.add(code)

            title = art["title"]
            category = art.get("category") or pending.get("category") or ""
            extra = parsed.get("extra") or ""
            product_name = clean_product_title(title or category)
            short_desc = rewrite_short_description(
                art.get("description") or "", product_name
            )
            diam = parsed.get("Diamètre") or ""
            prefix_src = last_prefix
            last_prefix = ""
            for src in (extra, prefix_src, line):
                if not src:
                    continue
                tok = src.split()[0]
                compact = tok.replace(" ", "")
                if DIAM_RE.match(compact) or DIAM_RE.match(tok):
                    diam = compact.replace("×", "X").replace("x", "X")
                    break
            if extra and (
                re.search(r"Ø|L\.\s*\d", extra, re.I)
                or DIM_LINE_RE.match(extra.replace(" ", ""))
                or DIAM_RE.match(extra.replace(" ", ""))
            ):
                diam = diam or extra.replace("×", "X").replace("x", "X")
            section = section_for_article(art, cy)
            variant_parts = []
            if section:
                variant_parts = [p for p in (section.get("group"), section.get("sub")) if p]
            elif art.get("variant_group") or art.get("variant_sub"):
                variant_parts = [
                    p for p in (art.get("variant_group"), art.get("variant_sub")) if p
                ]
            variant_label = " · ".join(variant_parts)
            if not variant_label:
                variant_label = extra or diam
            variant_label = re.sub(r"\s+\d+,\d{2}(\s+\d{5,})+.*$", "", variant_label or "").strip()
            variant_label = re.sub(r"\s+\d{6,}.*$", "", variant_label).strip()
            if re.search(r"circuit", extra or "", re.I):
                diam = ""
            _peeled, diam_from_title, tech_note = peel_technical(title or category)
            if diam_from_title and not diam:
                diam = diam_from_title
            note = clean_spaces(" · ".join(x for x in (art.get("note"), tech_note) if x))
            designation = product_name or build_designation(category, title, "")
            brand = detect_series_brand(designation) or art.get("brand") or ""
            if not brand and any(
                k in (title or "").upper()
                for k in ("COUDE", "COURBE", "RACCORD", "MANCHON", "MAMELON", "TÉ")
            ):
                brand = "Altech"
            if any(k in (category or "").upper() for k in ("RACCORDS CUIVRE", "RACCORDS LAITON")):
                if (brand or "").upper() in ("CUPROLIFE", "STARFIX", "WICU", "SANCO", "WIELAND", "ALTECH"):
                    if not any(
                        k in (title or "").upper()
                        for k in ("STARFIX", "CUPROLIFE", "WICU", "SANCO")
                    ):
                        brand = "Altech"

            img = art.get("image") or ""
            img_path = f"{IMAGE_REL}/{img}" if img else ""
            photo_letter = parsed.get("letter") or ""

            products.append(
                {
                    "Réf.Pro": parsed["Réf.Pro"],
                    "Réf.Four": parsed["Réf.Four"],
                    "Diamètre": diam,
                    "Vendu par": parsed["Vendu par"],
                    "Marque": brand,
                    "Désignation": designation,
                    "Variante": variant_label,
                    "Ordre": len(products) + 1,
                    "Code": code,
                    "Code interne": None,
                    "Famille": infer_famille(category, title),
                    "Catégorie": category,
                    "Prix HT": parsed["Prix HT"],
                    "Note": note,
                    "Description": short_desc,
                    "Image": img_path,
                    "Image produit": img_path,
                    "_photo_letter": photo_letter,
                    "_page": page_no,
                    "_art_y0": art.get("y0"),
                    "_art_col": 0 if art.get("col") == "L" else 1,
                }
            )

        left_ys, right_ys = red_separator_ys(fitz_doc[idx], pw, ph)
        title_pts = []
        seen_ty = set()
        for item in products[page_start:]:
            ty = item.get("_art_y0")
            col = item.get("_art_col")
            tx = None
            if ty is None or col is None:
                ty, tx = find_title_y(tp, (item.get("Désignation") or "")[:28], ph)
                col = 0 if (tx or 0) < pw / 2 else 1
            else:
                tx = 40.0 if int(col) == 0 else pw * 0.75
            if ty is None:
                continue
            key = (round(float(ty), 0), int(col))
            if key in seen_ty:
                continue
            seen_ty.add(key)
            title_pts.append((float(ty), float(tx or 0)))
        for ty, tx in title_pts:
            if tx < pw / 2:
                left_ys.append(ty)
            else:
                right_ys.append(ty)
        left_ys = sorted(set(round(y, 1) for y in left_ys))
        right_ys = sorted(set(round(y, 1) for y in right_ys))

        assign_block_photos(
            products[page_start:],
            photos,
            fitz_doc[idx],
            tp,
            ph,
            pw,
            left_ys,
            right_ys,
        )
        for item in products[page_start:]:
            hero = item.get("_hero") or ""
            if hero:
                item["Image produit"] = f"{IMAGE_REL}/{hero}"
            elif item.get("Image"):
                item["Image produit"] = item["Image"]

        if (idx + 1) % 25 == 0:
            print(
                f"Page {idx + 1}/{len(pdf)} -> {len(products)} produits, "
                f"{len(hash_to_name)} images (natif PNG, personnes exclues={skip_stats['people']})",
                flush=True,
            )

    fitz_doc.close()
    if FILTER_FAMILLE and not os.environ.get("EXTRACT_NO_FAMILY_FILTER"):
        products = [p for p in products if (p.get("Famille") or "") == FILTER_FAMILLE]
        print(f"Filtre famille {FILTER_FAMILLE} -> {len(products)} lignes", flush=True)
    print(
        f"Terminé {len(pdf)} pages -> {len(products)} produits, "
        f"{len(hash_to_name)} images pièces, {skip_stats['people']} photos personnes ignorées",
        flush=True,
    )
    verify_extracted_products(products)
    return products


GOLDEN_SKU_CHECKS = [
    {
        "code": "7782955",
        "name_contains": "FIXOPLAC",
        "name_forbids": ["GLISSEMENT", "SERTIR"],
        "desc_contains": ["plaque", "vis"],
        "variant_contains": "GLISSEMENT",
        "brand": "Ayor",
    },
    {
        "code": "7782956",
        "name_contains": "FIXOPLAC",
        "name_forbids": ["GLISSEMENT"],
        "variant_contains": "SERTIR",
    },
    {
        "code": "7782959",
        "name_contains": "FIXOPLAC",
        "variant_contains": "SOUS",
    },
    {
        "code": "4841911",
        "name_contains": "SMART U",
        "name_forbids": ["FIXOPLAC"],
        "variant_contains": "GLISSEMENT",
    },
    {
        "code": "6301286",
        "name_contains": "STARFIX",
        "name_forbids": ["FIXOPLAC", "MONOTROU"],
    },
    {
        "code": "4011055",
        "name_contains": "PLENITUDE",
        "name_forbids": ["UNI"],
        "variant_contains": "UNI",
    },
]


def verify_extracted_products(products):
    """Contrôles : titre d'article ≠ sous-rubrique de tableau, SKU témoins du PDF."""
    by_code = {str(p.get("Code") or ""): p for p in products}
    errors = []
    warnings = []

    for spec in GOLDEN_SKU_CHECKS:
        code = spec["code"]
        row = by_code.get(code)
        if not row:
            warnings.append(f"{code}: absent de l'extraction (hors pages / hors famille)")
            continue
        name = (row.get("Désignation") or "").upper()
        desc = (row.get("Description") or "").lower()
        variant = (row.get("Variante") or "").upper()
        brand = row.get("Marque") or ""
        if spec.get("name_contains") and spec["name_contains"].upper() not in name:
            errors.append(
                f"{code}: nom {row.get('Désignation')!r} sans {spec['name_contains']}"
            )
        for forbid in spec.get("name_forbids") or []:
            if forbid.upper() in name:
                errors.append(f"{code}: nom {row.get('Désignation')!r} contient {forbid}")
        for needle in spec.get("desc_contains") or []:
            if needle.lower() not in desc:
                errors.append(f"{code}: description sans {needle!r} ({row.get('Description')!r})")
        if spec.get("variant_contains") and spec["variant_contains"].upper() not in variant:
            errors.append(
                f"{code}: variante {row.get('Variante')!r} sans {spec['variant_contains']}"
            )
        if spec.get("brand") and spec["brand"].lower() not in brand.lower():
            errors.append(f"{code}: marque {brand!r} ≠ {spec['brand']}")

    suspect = []
    hard = []
    for p in products:
        name = (p.get("Désignation") or "").strip()
        if TECH_VARIANT_RE.match(name) or name.upper() in {"UNI", "BICOLORE"}:
            hard.append(f"{p.get('Code')} {name!r}")
        elif is_table_section_header(name):
            suspect.append(f"{p.get('Code')} {name!r}")
    if suspect:
        warnings.append(
            f"{len(suspect)} titres encore type tableau (ex. {', '.join(suspect[:8])})"
        )
    if len(hard) >= 12:
        warnings.append(f"{len(hard)} titres UNI/variante technique (ex. {', '.join(hard[:8])})")
        errors.append(
            f"Trop de titres-tableau restants ({len(hard)}). Relire les traits verts / colonnes."
        )

    for msg in warnings:
        print(f"Vérif warning: {msg}")
    for msg in errors:
        print(f"Vérif ERREUR: {msg}")
    if errors and not os.environ.get("EXTRACT_SKIP_VERIFY"):
        raise SystemExit(f"Extraction rejetée: {len(errors)} contrôle(s) en échec")
    print(
        f"Vérif OK: {len(GOLDEN_SKU_CHECKS)} SKU témoins, "
        f"{len(warnings)} warning(s), {len(products)} lignes"
    )


def write_excel(products):
    products.sort(
        key=lambda p: (
            p.get("Famille") or "",
            p.get("Catégorie") or "",
            p.get("Désignation") or "",
            int(p.get("Ordre") or 0),
        )
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produits"
    fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c, name in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, name)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    for r, p in enumerate(products, 2):
        for c, name in enumerate(COLUMNS, 1):
            cell = ws.cell(r, c, sanitize_cell(p.get(name, "")))
            cell.alignment = Alignment(
                vertical="center", wrap_text=name in ("Désignation", "Variante", "Note", "Description")
            )
    widths = {
        "A": 12, "B": 16, "C": 16, "D": 12, "E": 18, "F": 42,
        "G": 22, "H": 8, "I": 12, "J": 14, "K": 20, "L": 28, "M": 12, "N": 26, "O": 42, "P": 36,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:Q{max(1, len(products) + 1)}"
    wb.save(OUTPUT_PATH)


def patch_existing_xlsx(products, xlsx_path):
    """Met à jour Désignation / Marque / Catégorie / Description des lignes déjà présentes."""
    by_code = {}
    for p in products:
        code = str(p.get("Code") or "").strip()
        if code:
            by_code[code] = p
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    if "Description" not in headers:
        note_idx = headers.index("Note") + 1 if "Note" in headers else len(headers)
        ws.insert_cols(note_idx + 1)
        ws.cell(1, note_idx + 1, "Description")
        headers = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(headers) if name}
    updated = 0
    for row in range(2, ws.max_row + 1):
        code_cell = ws.cell(row, col["Code"]).value
        if code_cell is None:
            continue
        code = str(code_cell).strip().replace(".0", "")
        src = by_code.get(code)
        if not src:
            continue
        if "Désignation" in col and src.get("Désignation"):
            ws.cell(row, col["Désignation"], sanitize_cell(src["Désignation"]))
        if "Marque" in col and src.get("Marque"):
            ws.cell(row, col["Marque"], sanitize_cell(src["Marque"]))
        if "Catégorie" in col and src.get("Catégorie"):
            ws.cell(row, col["Catégorie"], sanitize_cell(src["Catégorie"]))
        if "Description" in col:
            ws.cell(row, col["Description"], sanitize_cell(src.get("Description") or ""))
        if "Variante" in col:
            ws.cell(row, col["Variante"], sanitize_cell(src.get("Variante") or ""))
        if "Diamètre" in col:
            ws.cell(row, col["Diamètre"], sanitize_cell(src.get("Diamètre") or ""))
        if "Réf.Pro" in col and src.get("Réf.Pro"):
            ws.cell(row, col["Réf.Pro"], sanitize_cell(src["Réf.Pro"]))
        updated += 1
    wb.save(xlsx_path)
    print(f"Patch Excel: {updated} lignes mises à jour → {xlsx_path}")


def main():
    products = process_pdf()
    patch_path = os.environ.get("PATCH_EXISTING_XLSX")
    if patch_path:
        patch_existing_xlsx(products, patch_path)
    if os.environ.get("EXTRACT_DEBUG_NO_EXCEL"):
        print(f"DEBUG skip excel ({len(products)} rows)")
        return
    write_excel(products)
    priced = sum(1 for p in products if p.get("Prix HT"))
    imaged = sum(1 for p in products if p.get("Image"))
    print(f"\nExporté {len(products)} produits -> {OUTPUT_PATH}")
    print(f"Dossier images -> {IMAGE_DIR}")
    print(f"Avec prix: {priced}/{len(products)} ({100 * priced / max(1, len(products)):.1f}%)")
    print(f"Avec image: {imaged}/{len(products)} ({100 * imaged / max(1, len(products)):.1f}%)")
    print("Familles:", Counter(p.get("Famille") or "?" for p in products).most_common())


if __name__ == "__main__":
    main()
