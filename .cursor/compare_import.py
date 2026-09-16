#!/usr/bin/env python3
"""Compare le nouvel Excel d'extraction à la base, famille Plomberie."""
import re
import subprocess
from collections import Counter, defaultdict

import openpyxl

XLSX = "/Users/gael/Desktop/catalogue/catalogue_pro_2026_plomberie.xlsx"
DB = "postgresql://gael@localhost:5432/afecon_catalogue"


def psql(sql):
    out = subprocess.check_output(["psql", DB, "-At", "-F", "\t", "-c", sql], text=True)
    return [l.split("\t") for l in out.splitlines() if l.strip()]


def norm(s):
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def suspect(name):
    n = (name or "").strip()
    if not n:
        return "vide"
    if re.match(r"^[A-Z]\s+\S", n):
        return "prefixe_lettre"
    if " " not in n and re.search(r"\d", n) and re.search(r"[A-Za-z]", n):
        return "ref_fournisseur"
    if n[:1].islower():
        return "minuscule"
    return None


# --- Excel ---------------------------------------------------------------
wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ws = wb.active
head = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
ix = {h: i for i, h in enumerate(head)}
xl = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    code = str(row[ix["Code"]] or "").strip()
    if not code:
        continue
    xl[code] = {
        "name": str(row[ix["Désignation"]] or "").strip(),
        "variant": str(row[ix["Variante"]] or "").strip(),
        "image": str(row[ix["Image"]] or "").strip(),
    }
wb.close()

# --- Base ----------------------------------------------------------------
rows = psql("""
    SELECT r.code, p.id, p.name, COALESCE(r.variant_label,''), COALESCE(p.image_path,'')
    FROM references_sku r
    JOIN products p ON p.id = r.product_id
    JOIN categories c ON c.id = p.category_id
    WHERE c.family_id = 1
""")
db = {c: {"pid": int(pid), "name": n, "variant": v, "image": img} for c, pid, n, v, img in rows}

# --- Couverture ----------------------------------------------------------
common = sorted(set(db) & set(xl))
only_db = sorted(set(db) - set(xl))
only_xl = sorted(set(xl) - set(db))

print("=" * 74)
print("COUVERTURE")
print(f"  references Plomberie en base   : {len(db)}")
print(f"  lignes dans le nouvel Excel    : {len(xl)}")
print(f"  codes communs                  : {len(common)}")
print(f"  en base mais PAS dans l'Excel  : {len(only_db)}   <- ne seront pas corriges")
print(f"  nouveaux codes (Excel seul)    : {len(only_xl)}")

# --- Noms d'article ------------------------------------------------------
name_same = name_fix = name_reg = name_other = 0
fixes, regressions, others = [], [], []
for c in common:
    dn, xn = db[c]["name"], xl[c]["name"]
    if norm(dn) == norm(xn):
        name_same += 1
        continue
    sd, sx = suspect(dn), suspect(xn)
    if sd and not sx:
        name_fix += 1
        fixes.append((c, dn, xn, sd))
    elif sx and not sd:
        name_reg += 1
        regressions.append((c, dn, xn, sx))
    else:
        name_other += 1
        others.append((c, dn, xn))

print()
print("=" * 74)
print(f"NOM D'ARTICLE (sur {len(common)} codes communs)")
print(f"  identique                      : {name_same}")
print(f"  CORRIGE (suspect -> propre)    : {name_fix}")
print(f"  REGRESSION (propre -> suspect) : {name_reg}")
print(f"  autre changement               : {name_other}")

if fixes:
    print("\n  exemples de corrections :")
    for c, dn, xn, sd in fixes[:12]:
        print(f"    {c}  [{sd:15}] {dn[:30]:30} -> {xn[:38]}")
if regressions:
    print("\n  REGRESSIONS a examiner :")
    for c, dn, xn, sx in regressions[:20]:
        print(f"    {c}  [{sx:15}] {dn[:30]:30} -> {xn[:38]}")
if others:
    print("\n  autres changements (echantillon) :")
    for c, dn, xn in others[:15]:
        print(f"    {c}  {dn[:33]:33} -> {xn[:36]}")

# --- Variantes -----------------------------------------------------------
v_same = v_gained = v_lost = v_changed = 0
lost = []
for c in common:
    dv, xv = db[c]["variant"], xl[c]["variant"]
    if norm(dv) == norm(xv):
        v_same += 1
    elif not dv and xv:
        v_gained += 1
    elif dv and not xv:
        v_lost += 1
        lost.append((c, db[c]["name"], dv))
    else:
        v_changed += 1

print()
print("=" * 74)
print(f"VARIANTE (sur {len(common)} codes communs)")
print(f"  identique          : {v_same}")
print(f"  GAGNEE (vide -> x) : {v_gained}")
print(f"  PERDUE (x -> vide) : {v_lost}")
print(f"  modifiee           : {v_changed}")
if lost:
    print("\n  variantes perdues (echantillon) :")
    for c, n, dv in lost[:15]:
        print(f"    {c}  {n[:32]:32} perd '{dv[:30]}'")

# --- Photos --------------------------------------------------------------
img_same = img_changed = img_none = 0
for c in common:
    xi = xl[c]["image"].replace("\\", "/").split("/")[-1].lower()
    di = db[c]["image"].split("/")[-1].lower()
    if not xi:
        img_none += 1
    elif xi == di:
        img_same += 1
    else:
        img_changed += 1
print()
print("=" * 74)
print("PHOTO")
print(f"  meme fichier    : {img_same}")
print(f"  fichier different : {img_changed}")
print(f"  aucune photo dans l'Excel : {img_none}")

# --- Prevision orphelins -------------------------------------------------
prod_refs = defaultdict(list)
for c in db:
    prod_refs[db[c]["pid"]].append(c)

pid_name = {db[c]["pid"]: db[c]["name"] for c in db}
orphans = []
for pid, codes in prod_refs.items():
    kept = [c for c in codes if c not in xl or norm(xl[c]["name"]) == norm(pid_name[pid])]
    if not kept:
        orphans.append((pid, pid_name[pid], len(codes)))

print()
print("=" * 74)
print(f"PREVISION ORPHELINS apres import : {len(orphans)} produits perdraient toutes leurs references")
by_kind = Counter(suspect(n) or "nom_propre" for _p, n, _k in orphans)
for k, v in by_kind.most_common():
    print(f"    {k:18} {v}")
print("\n  echantillon :")
for pid, n, k in orphans[:20]:
    print(f"    id={pid:<7} {k} ref(s)  {n[:48]}")
