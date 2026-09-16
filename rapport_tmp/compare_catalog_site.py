#!/usr/bin/env python3
"""Compare catalogue Excel vs site Afecon — refs, noms, variantes, photos (produit vs marque)."""
from __future__ import annotations
import csv, hashlib, json, os, re, sys, time, urllib.request
from collections import Counter, defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl --quiet")
    import openpyxl

try:
    from PIL import Image
except ImportError:
    os.system(f"{sys.executable} -m pip install pillow --quiet")
    from PIL import Image

OUT_DIR = Path("/Users/gael/Documents/htdocs/AfeconCatalogue/rapport_tmp")
OUT_DIR.mkdir(parents=True, exist_ok=True)
REPORT_MD = Path("/Users/gael/Documents/htdocs/AfeconCatalogue/rapport_ecarts_catalogue_site.md")
REPORT_CSV = Path("/Users/gael/Documents/htdocs/AfeconCatalogue/rapport_ecarts_catalogue_site.csv")
API = "http://localhost:4000"
UPLOADS = Path("/Users/gael/Documents/htdocs/AfeconCatalogue/backend/uploads")

EXCELS = [
    Path("/Users/gael/Desktop/Catalogue/catalogue_pro_2026_plomberie.xlsx"),
    Path("/Users/gael/Desktop/catalogue_pro_2026_hq.xlsx"),
    Path("/Users/gael/Desktop/Catalogue/afeconCatalogue_basededonnee -copie.xlsx"),
    Path("/Users/gael/Desktop/Catalogue/electricite_legrand.xlsx"),
    Path("/Users/gael/Desktop/catalogue_pro_2026_plomberie.xlsx"),
]

IMAGE_DIRS = [
    Path("/Users/gael/Desktop/catalogue_pro_2026_images"),
    Path("/Users/gael/Desktop/catalogue_pro_2026_images_hq"),
    Path("/Users/gael/Desktop/Catalogue/images_hq"),
    Path("/Users/gael/Desktop/Catalogue/images_fix_merge"),
    Path("/Users/gael/Desktop/catalogue/images_hq"),
]

def log(msg):
    print(msg, flush=True)

def norm(s):
    if s is None:
        return ""
    s = str(s).strip().upper()
    s = re.sub(r"\s+", " ", s)
    return s

def stem(path_or_name):
    if not path_or_name:
        return ""
    name = os.path.basename(str(path_or_name).replace("\\", "/"))
    return os.path.splitext(name)[0].lower()

def file_md5(p: Path, limit=2_000_000):
    h = hashlib.md5()
    try:
        with open(p, "rb") as f:
            while True:
                chunk = f.read(65536)
                if not chunk:
                    break
                h.update(chunk)
                if f.tell() > limit and limit:
                    # full hash for accuracy on photos — remove limit for exact
                    pass
        # re-hash full file for exact match
    except Exception:
        return None
    h = hashlib.md5()
    try:
        with open(p, "rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return None

def classify_image(path: Path, usage_count: int):
    """Heuristique: produit utile vs marque/marketing/inutile."""
    reasons = []
    kind = "produit"
    try:
        with Image.open(path) as im:
            w, h = im.size
            ratio = w / h if h else 0
            # very small
            if w < 120 or h < 120:
                kind = "marque_ou_inutile"
                reasons.append(f"petite({w}x{h})")
            # extreme banner ratio
            if ratio > 3.5 or (ratio and ratio < 0.28):
                kind = "marque_ou_inutile"
                reasons.append(f"bandeau(ratio={ratio:.2f})")
            # reused on many SKUs
            if usage_count >= 25:
                kind = "marque_ou_inutile"
                reasons.append(f"réutilisée_x{usage_count}")
            # mostly white / low variance heuristic via thumbnail
            im2 = im.convert("RGB").resize((32, 32))
            pixels = list(im2.getdata())
            avg = sum(sum(p) for p in pixels) / (len(pixels) * 3)
            uniq = len(set(pixels))
            if uniq < 8 and avg > 230:
                kind = "marque_ou_inutile"
                reasons.append("quasi_vide_blanc")
            if uniq < 12 and avg < 40:
                kind = "marque_ou_inutile"
                reasons.append("quasi_noir")
    except Exception as e:
        reasons.append(f"err:{e}")
        kind = "inconnu"
    return kind, reasons

def fetch_json(url):
    with urllib.request.urlopen(url, timeout=60) as r:
        return json.load(r)

def dump_site():
    cache = OUT_DIR / "site_products_full.json"
    if cache.exists() and cache.stat().st_mtime > time.time() - 3600:
        log(f"Cache site: {cache}")
        return json.loads(cache.read_text())
    log("Export site via API…")
    page = 1
    limit = 100
    items = []
    total = None
    while True:
        data = fetch_json(f"{API}/api/products?limit={limit}&page={page}")
        total = data.get("total", total)
        batch = data.get("items") or []
        if not batch:
            break
        items.extend(batch)
        log(f"  list page {page}: {len(items)}/{total}")
        if len(items) >= (total or 0):
            break
        page += 1
        if page > 500:
            break
    # enrich with references
    full = []
    for i, it in enumerate(items, 1):
        try:
            detail = fetch_json(f"{API}/api/products/{it['id']}")
        except Exception as e:
            detail = dict(it)
            detail["references"] = []
            detail["_err"] = str(e)
        full.append(detail)
        if i % 100 == 0:
            log(f"  detail {i}/{len(items)}")
    cache.write_text(json.dumps(full, ensure_ascii=False))
    log(f"Site dump: {len(full)} products")
    return full

def load_excels():
    rows = []
    seen = set()
    for path in EXCELS:
        if not path.exists():
            log(f"Excel manquant: {path}")
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            it = ws.iter_rows(values_only=True)
            try:
                header = next(it)
            except StopIteration:
                continue
            header = [str(h).strip() if h is not None else "" for h in header]
            # map columns flexibly
            idx = {h: i for i, h in enumerate(header)}
            def col(*names):
                for n in names:
                    if n in idx:
                        return idx[n]
                # case-insensitive
                low = {h.lower(): i for i, h in enumerate(header)}
                for n in names:
                    if n.lower() in low:
                        return low[n.lower()]
                return None
            i_code = col("Code", "code")
            i_refpro = col("Réf.Pro", "Ref.Pro", "Réf Pro")
            i_reffour = col("Réf.Four", "Ref.Four")
            i_marque = col("Marque")
            i_des = col("Désignation", "Designation", "Nom")
            i_var = col("Variante", "Variant")
            i_fam = col("Famille")
            i_cat = col("Catégorie", "Categorie")
            i_img = col("Image", "image", "image_path")
            n = 0
            for r in it:
                if not r:
                    continue
                def g(i):
                    return r[i] if i is not None and i < len(r) else None
                code = g(i_code)
                refpro = g(i_refpro)
                des = g(i_des)
                if code is None and refpro is None and not des:
                    continue
                key = (norm(code), norm(refpro), norm(des), norm(g(i_var)), str(path))
                if key in seen:
                    continue
                seen.add(key)
                rows.append({
                    "source": path.name,
                    "code": str(code).strip() if code is not None else "",
                    "ref_pro": str(refpro).strip() if refpro is not None else "",
                    "ref_four": str(g(i_reffour)).strip() if g(i_reffour) is not None else "",
                    "marque": str(g(i_marque)).strip() if g(i_marque) is not None else "",
                    "designation": str(des).strip() if des is not None else "",
                    "variante": str(g(i_var)).strip() if g(i_var) is not None else "",
                    "famille": str(g(i_fam)).strip() if g(i_fam) is not None else "",
                    "categorie": str(g(i_cat)).strip() if g(i_cat) is not None else "",
                    "image": str(g(i_img)).strip() if g(i_img) is not None else "",
                })
                n += 1
            log(f"Excel {path.name} / {sheet}: {n} lignes")
        wb.close()
    return rows

def resolve_catalog_image(img_field: str):
    if not img_field:
        return None
    p = Path(img_field)
    if p.is_file():
        return p
    # relative names
    name = os.path.basename(img_field)
    for d in IMAGE_DIRS:
        cand = d / name
        if cand.is_file():
            return cand
        # maybe path like catalogue_pro_2026_images/p0907_02.jpg
        parts = Path(img_field).parts
        if len(parts) >= 2:
            cand2 = d / parts[-1]
            if cand2.is_file():
                return cand2
    # search by stem across dirs (first hit)
    st = stem(img_field)
    for d in IMAGE_DIRS:
        if not d.is_dir():
            continue
        for ext in (".jpg", ".jpeg", ".png", ".webp"):
            cand = d / f"{st}{ext}"
            if cand.is_file():
                return cand
    return None

def resolve_site_image(img_path: str):
    if not img_path:
        return None
    name = os.path.basename(img_path)
    cand = UPLOADS / name
    if cand.is_file():
        return cand
    # absolute-ish
    p = Path(img_path)
    if p.is_file():
        return p
    return None

def main():
    t0 = time.time()
    site = dump_site()
    catalog = load_excels()
    log(f"Catalogue lignes: {len(catalog)} | Site produits: {len(site)}")

    # index site by code / ref_pro / supplier
    site_by_code = {}
    site_by_refpro = {}
    site_by_supplier = {}
    site_refs = []
    image_usage_site = Counter()
    for p in site:
        refs = p.get("references") or []
        for ref in refs:
            site_refs.append((p, ref))
            code = str(ref.get("supplier_code") or ref.get("code") or "").strip()
            # supplier_code often = Code catalogue; also code AFE
            sc = str(ref.get("supplier_code") or "").strip()
            afe = str(ref.get("code") or "").strip()
            rp = str(ref.get("ref_pro") or "").strip()
            rf = str(ref.get("ref_four") or "").strip()
            if sc:
                site_by_code[norm(sc)] = (p, ref)
            if afe:
                site_by_code[norm(afe)] = (p, ref)
            if rp:
                site_by_refpro[norm(rp)] = (p, ref)
            if rf:
                site_by_supplier[norm(rf)] = (p, ref)
            img = ref.get("image_path") or p.get("image_path") or p.get("display_image")
            if img:
                image_usage_site[stem(img)] += 1

    # catalog image usage
    image_usage_cat = Counter()
    for row in catalog:
        if row["image"]:
            image_usage_cat[stem(row["image"])] += 1

    # pre-index catalog images by stem across dirs
    catalog_files_by_stem = {}
    for d in IMAGE_DIRS:
        if not d.is_dir():
            continue
        for f in d.iterdir():
            if f.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
                catalog_files_by_stem.setdefault(f.stem.lower(), f)

    site_files_by_stem = {}
    if UPLOADS.is_dir():
        for f in UPLOADS.iterdir():
            if f.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
                site_files_by_stem.setdefault(f.stem.lower(), f)

    gaps = []
    stats = Counter()

    def add_gap(typ, **kw):
        stats[typ] += 1
        gaps.append({"type_ecart": typ, **kw})

    matched_codes = set()

    for row in catalog:
        code_n = norm(row["code"])
        rp_n = norm(row["ref_pro"])
        rf_n = norm(row["ref_four"])
        hit = None
        if code_n and code_n in site_by_code:
            hit = site_by_code[code_n]
        elif rp_n and rp_n in site_by_refpro:
            hit = site_by_refpro[rp_n]
        elif rf_n and rf_n in site_by_supplier:
            hit = site_by_supplier[rf_n]

        if not hit:
            add_gap(
                "ref_absente_site",
                code=row["code"],
                ref_pro=row["ref_pro"],
                designation_catalogue=row["designation"],
                designation_site="",
                variante_cat=row["variante"],
                variante_site="",
                image_cat=row["image"],
                image_site="",
                detail=f"source={row['source']}",
            )
            continue

        p, ref = hit
        matched_codes.add(id(ref))
        # name compare: product name vs designation — often product is short group name
        site_name = norm(p.get("name"))
        site_var = norm(ref.get("variant_label"))
        cat_des = norm(row["designation"])
        cat_var = norm(row["variante"])
        # designation match soft: exact or containment
        if cat_des and site_name and cat_des != site_name and cat_des not in site_name and site_name not in cat_des:
            # also compare to variant
            if cat_des != site_var:
                add_gap(
                    "nom_different",
                    code=row["code"],
                    ref_pro=row["ref_pro"],
                    designation_catalogue=row["designation"],
                    designation_site=p.get("name") or "",
                    variante_cat=row["variante"],
                    variante_site=ref.get("variant_label") or "",
                    image_cat=row["image"],
                    image_site=ref.get("image_path") or p.get("image_path") or "",
                    detail="désignation ≠ nom produit site",
                )
            else:
                stats["nom_ok_via_variante"] += 1
        else:
            stats["nom_ok"] += 1

        if cat_var and site_var and cat_var != site_var and cat_var not in site_var and site_var not in cat_var:
            add_gap(
                "variante_differente",
                code=row["code"],
                ref_pro=row["ref_pro"],
                designation_catalogue=row["designation"],
                designation_site=p.get("name") or "",
                variante_cat=row["variante"],
                variante_site=ref.get("variant_label") or "",
                image_cat=row["image"],
                image_site=ref.get("image_path") or p.get("image_path") or "",
                detail="",
            )
        elif cat_var and not site_var:
            add_gap(
                "variante_manquante_site",
                code=row["code"],
                ref_pro=row["ref_pro"],
                designation_catalogue=row["designation"],
                designation_site=p.get("name") or "",
                variante_cat=row["variante"],
                variante_site="",
                image_cat=row["image"],
                image_site=ref.get("image_path") or p.get("image_path") or "",
                detail="",
            )
        else:
            stats["variante_ok"] += 1

        # PHOTOS
        site_img_field = ref.get("image_path") or p.get("image_path") or p.get("display_image") or ""
        cat_img_path = resolve_catalog_image(row["image"]) if row["image"] else catalog_files_by_stem.get(stem(row["image"]))
        # if excel has no image, try matching by common stem later — skip
        site_img_path = resolve_site_image(site_img_field)

        usage = max(image_usage_cat.get(stem(row["image"]), 0), image_usage_site.get(stem(site_img_field), 0))
        kind_cat, reasons_cat = ("inconnu", [])
        kind_site, reasons_site = ("inconnu", [])
        if cat_img_path and cat_img_path.is_file():
            kind_cat, reasons_cat = classify_image(cat_img_path, image_usage_cat.get(stem(cat_img_path.name), usage))
        if site_img_path and site_img_path.is_file():
            kind_site, reasons_site = classify_image(site_img_path, image_usage_site.get(stem(site_img_path.name), usage))

        # if either side brand/useless — track separately
        if kind_cat == "marque_ou_inutile" or kind_site == "marque_ou_inutile":
            add_gap(
                "photo_marque_ou_inutile",
                code=row["code"],
                ref_pro=row["ref_pro"],
                designation_catalogue=row["designation"],
                designation_site=p.get("name") or "",
                variante_cat=row["variante"],
                variante_site=ref.get("variant_label") or "",
                image_cat=row["image"] or (cat_img_path.name if cat_img_path else ""),
                image_site=site_img_field,
                detail=f"cat={kind_cat}:{','.join(reasons_cat)}; site={kind_site}:{','.join(reasons_site)}",
            )
            # still check match but not as critical product mismatch
        is_product_photo = kind_cat != "marque_ou_inutile" and kind_site != "marque_ou_inutile"

        if not site_img_field or not site_img_path:
            add_gap(
                "photo_manquante_site" if is_product_photo else "photo_manquante_site_marque",
                code=row["code"],
                ref_pro=row["ref_pro"],
                designation_catalogue=row["designation"],
                designation_site=p.get("name") or "",
                variante_cat=row["variante"],
                variante_site=ref.get("variant_label") or "",
                image_cat=row["image"],
                image_site=site_img_field,
                detail="produit" if is_product_photo else "marque/inutile",
            )
        elif row["image"] and not cat_img_path:
            add_gap(
                "photo_catalogue_introuvable_disque",
                code=row["code"],
                ref_pro=row["ref_pro"],
                designation_catalogue=row["designation"],
                designation_site=p.get("name") or "",
                variante_cat=row["variante"],
                variante_site=ref.get("variant_label") or "",
                image_cat=row["image"],
                image_site=site_img_field,
                detail="",
            )
        elif cat_img_path and site_img_path:
            same_stem = stem(cat_img_path.name) == stem(site_img_path.name)
            h1, h2 = file_md5(cat_img_path), file_md5(site_img_path)
            same_hash = h1 and h2 and h1 == h2
            if same_hash:
                stats["photo_identique_hash"] += 1
                if is_product_photo:
                    stats["photo_produit_ok"] += 1
                else:
                    stats["photo_marque_ok"] += 1
            elif same_stem:
                stats["photo_meme_stem_hash_diff"] += 1
                add_gap(
                    "photo_meme_nom_contenu_different" if is_product_photo else "photo_marque_meme_nom_contenu_diff",
                    code=row["code"],
                    ref_pro=row["ref_pro"],
                    designation_catalogue=row["designation"],
                    designation_site=p.get("name") or "",
                    variante_cat=row["variante"],
                    variante_site=ref.get("variant_label") or "",
                    image_cat=str(cat_img_path),
                    image_site=str(site_img_path),
                    detail=f"md5_cat={h1} md5_site={h2}",
                )
            else:
                add_gap(
                    "photo_differente" if is_product_photo else "photo_marque_differente",
                    code=row["code"],
                    ref_pro=row["ref_pro"],
                    designation_catalogue=row["designation"],
                    designation_site=p.get("name") or "",
                    variante_cat=row["variante"],
                    variante_site=ref.get("variant_label") or "",
                    image_cat=str(cat_img_path.name),
                    image_site=str(site_img_path.name),
                    detail=f"stem_cat={stem(cat_img_path.name)} stem_site={stem(site_img_path.name)}",
                )
        elif not row["image"] and site_img_path:
            stats["photo_site_sans_image_excel"] += 1

    # site refs not in catalog
    cat_codes = {norm(r["code"]) for r in catalog if r["code"]}
    cat_refpros = {norm(r["ref_pro"]) for r in catalog if r["ref_pro"]}
    cat_reffours = {norm(r["ref_four"]) for r in catalog if r["ref_four"]}
    for p, ref in site_refs:
        sc = norm(ref.get("supplier_code"))
        rp = norm(ref.get("ref_pro"))
        rf = norm(ref.get("ref_four"))
        afe = norm(ref.get("code"))
        if (sc and sc in cat_codes) or (rp and rp in cat_refpros) or (rf and rf in cat_reffours) or (afe and afe in cat_codes):
            stats["ref_site_trouvee_catalogue"] += 1
            continue
        add_gap(
            "ref_absente_catalogue",
            code=ref.get("supplier_code") or ref.get("code") or "",
            ref_pro=ref.get("ref_pro") or "",
            designation_catalogue="",
            designation_site=p.get("name") or "",
            variante_cat="",
            variante_site=ref.get("variant_label") or "",
            image_cat="",
            image_site=ref.get("image_path") or p.get("image_path") or "",
            detail=f"afe={ref.get('code')}",
        )

    # write CSV
    fields = [
        "type_ecart","code","ref_pro","designation_catalogue","designation_site",
        "variante_cat","variante_site","image_cat","image_site","detail",
    ]
    with REPORT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for g in gaps:
            w.writerow({k: g.get(k, "") for k in fields})

    # summary MD
    elapsed = time.time() - t0
    lines = []
    lines.append("# Rapport d’écarts catalogue ↔ site Afecon")
    lines.append("")
    lines.append(f"Généré en {elapsed:.0f}s.")
    lines.append("")
    lines.append("## Totaux")
    lines.append(f"- Produits site: **{len(site)}**")
    lines.append(f"- Références site: **{len(site_refs)}**")
    lines.append(f"- Lignes catalogue (Excel): **{len(catalog)}**")
    lines.append(f"- Fichiers uploads site: **{len(site_files_by_stem)}** stems")
    lines.append("")
    lines.append("## Stats")
    for k, v in sorted(stats.items(), key=lambda x: (-x[1], x[0])):
        lines.append(f"- `{k}`: **{v}**")
    lines.append("")
    lines.append("## Écarts (compteurs)")
    type_counts = Counter(g["type_ecart"] for g in gaps)
    for k, v in type_counts.most_common():
        lines.append(f"- **{k}**: {v}")
    lines.append("")
    lines.append("## Photos — distinction produit vs marque/inutile")
    lines.append("- `photo_produit_ok` / `photo_identique_hash`: match exact utile")
    lines.append("- `photo_differente` / `photo_meme_nom_contenu_different` / `photo_manquante_site`: **critiques** (vraies photos produit)")
    lines.append("- `photo_marque_ou_inutile` et types `*_marque*`: signalés à part (logo, bandeau, image très réutilisée, etc.)")
    lines.append("")
    lines.append("## Exemples critiques — photos produit")
    crit = [g for g in gaps if g["type_ecart"] in {
        "photo_differente","photo_meme_nom_contenu_different","photo_manquante_site"
    }][:15]
    if not crit:
        lines.append("_Aucun écart photo produit critique trouvé (ou pas assez de liens Image Excel)._")
    for g in crit:
        lines.append(f"- **{g['type_ecart']}** code=`{g.get('code')}` — cat `{g.get('designation_catalogue')}` / site `{g.get('designation_site')}` — `{g.get('image_cat')}` vs `{g.get('image_site')}` — {g.get('detail')}")
    lines.append("")
    lines.append("## Exemples — photos marque / inutiles")
    brand = [g for g in gaps if "marque" in g["type_ecart"]][:10]
    for g in brand:
        lines.append(f"- **{g['type_ecart']}** `{g.get('image_cat')}` / `{g.get('image_site')}` — {g.get('detail')}")
    lines.append("")
    lines.append("## Exemples — refs / noms / variantes")
    for typ in ("ref_absente_site","nom_different","variante_differente","ref_absente_catalogue"):
        ex = [g for g in gaps if g["type_ecart"]==typ][:5]
        if not ex:
            continue
        lines.append(f"### {typ}")
        for g in ex:
            lines.append(f"- code=`{g.get('code')}` ref_pro=`{g.get('ref_pro')}` | `{g.get('designation_catalogue')}` ↔ `{g.get('designation_site')}` | var `{g.get('variante_cat')}` ↔ `{g.get('variante_site')}`")
    lines.append("")
    lines.append(f"CSV détaillé: `{REPORT_CSV}`")
    lines.append("")
    lines.append("## Méthode")
    lines.append("- Catalogue = Excel pro 2026 (PDF trop volumineux pour transfert direct; Excel = extraction structurée).")
    lines.append("- Site = API locale `/api/products` + détails références.")
    lines.append("- Photos = comparaison stem fichier + hash MD5; classification marque/inutile par taille, ratio bandeau, réutilisation massive, image quasi vide.")
    REPORT_MD.write_text("\n".join(lines), encoding="utf-8")
    (OUT_DIR / "stats.json").write_text(json.dumps({"stats":dict(stats),"types":dict(type_counts),"n_gaps":len(gaps),"n_site":len(site),"n_refs":len(site_refs),"n_cat":len(catalog)}, ensure_ascii=False, indent=2), encoding="utf-8")
    log(f"DONE gaps={len(gaps)} md={REPORT_MD} csv={REPORT_CSV}")

if __name__ == "__main__":
    main()
